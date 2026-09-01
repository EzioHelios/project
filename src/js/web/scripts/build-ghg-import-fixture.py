from __future__ import annotations

# ruff: noqa: INP001
import argparse
import os
import sys
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from app.domain.ghg_inventory.infrastructure.workbook import build_template_bytes

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook


def _codes(workbook: Workbook, sheet_name: str) -> list[str]:
    return [str(cell.value or "") for cell in workbook[sheet_name][2]]


def _append(workbook: Workbook, sheet_name: str, payload: dict[str, object]) -> None:
    codes = _codes(workbook, sheet_name)
    workbook[sheet_name].append([payload.get(code) for code in codes])


def build_fixture(output: Path) -> None:
    suffix = datetime.now(UTC).strftime("%Y%m%d%H%M%S%f")
    prefix = os.environ.get("GHG_FIXTURE_PREFIX", "DEMO-PARK")
    area_code = os.environ.get("GHG_FIXTURE_AREA_CODE", "UTILITY")
    equipment_code = os.environ.get("GHG_FIXTURE_EQUIPMENT_CODE", "BOILER-01")
    equipment_name = os.environ.get("GHG_FIXTURE_EQUIPMENT_NAME", "燃气热水锅炉")
    meter_point = os.environ.get("GHG_FIXTURE_METER_POINT", "GAS-METER-01")
    context = {
        "inventory_code": os.environ.get("GHG_FIXTURE_INVENTORY_CODE", "DEMO-PARK-RP-2026"),
        "organization_code": os.environ.get("GHG_FIXTURE_ORGANIZATION_CODE", "DEMO-PARK-ORG"),
        "facility_code": os.environ.get("GHG_FIXTURE_FACILITY_CODE", "ROOT"),
        "iso_category_code": os.environ.get("GHG_FIXTURE_CATEGORY_CODE", "1.1"),
        "reporting_profile_code": os.environ.get("GHG_FIXTURE_REPORTING_PROFILE_CODE", "CORPORATE_ISO14064_GHGP"),
        "method_code": os.environ.get("GHG_FIXTURE_METHOD_CODE", "F01_NCV"),
        "area_codes": [code for code in os.environ.get("GHG_FIXTURE_AREA_CODES", f"ROOT,{area_code}").split(",") if code],
        "equipment_codes": [code for code in os.environ.get("GHG_FIXTURE_EQUIPMENT_CODES", equipment_code).split(",") if code],
        "measurement_boundary_codes": [],
    }
    workbook = load_workbook(BytesIO(build_template_bytes(context)))
    record_code = f"{prefix}-IMPORT-{suffix}-C11-GAS"
    source_code = f"{prefix}-IMPORT-{suffix}-SRC-GAS"
    parameter_set = f"{prefix}-IMPORT-{suffix}-PARAM-GAS"
    evidence_code = f"{prefix}-IMPORT-{suffix}-EVD-GAS"
    evidence_hash = sha256(f"demo import evidence {suffix}".encode()).hexdigest()

    boundary = workbook["01_组织与边界"]
    for column, code in enumerate(_codes(workbook, "01_组织与边界"), 1):
        if code == "boundary_method":
            boundary.cell(row=3, column=column, value="operational_control")

    _append(
        workbook,
        "02_排放源",
        {
            "source_code": source_code,
            "source_name": "演示导入固定源天然气燃烧",
            "activity_type": "stationary_combustion",
            "facility_code": context["facility_code"],
            "iso_category_code": context["iso_category_code"],
        },
    )
    _append(
        workbook,
        "03_活动数据",
        {
            "record_code": record_code,
            "source_code": source_code,
            "activity_code": "stationary_gas_combustion",
            "activity_name": "演示导入固定源天然气燃烧",
            "period_start": "2026-08-01",
            "period_end": "2026-08-31",
            "activity_amount": 1.25,
            "activity_unit": "万Nm3",
            "data_source_type": "measured",
            "measurement_method": "演示流量计月度读数",
            "parameter_set_code": parameter_set,
            "national_source_category": "固定燃烧",
            "boundary_status": "included",
            "remarks": "组织碳截图证据自动导入样例",
            "facility_or_device": equipment_name,
            "fuel_or_material": "天然气",
            "meter_point": meter_point,
        },
    )
    _append(
        workbook,
        "03A_区域设备归集",
        {
            "allocation_code": f"{record_code}-ALLOC-01",
            "record_code": record_code,
            "source_code": source_code,
            "area_code": area_code,
            "equipment_code": equipment_code,
            "allocation_role": "physical_origin",
            "allocation_ratio": 1,
            "source_kind": "manual",
            "valid_from": "2026-08-01",
            "valid_to": "2026-08-31",
        },
    )

    parameters = {
        "NCV": ("低位发热量", "method_parameter", 389.31, "GJ/万Nm3", "authority_published", "GB/T 2589-2020 综合能耗计算通则", "GB/T 2589-2020"),
        "CARBON_CONTENT": ("单位热值含碳量", "method_parameter", 0.0153, "tC/GJ", "authority_published", "省级温室气体清单指南演示因子", "企业演示假设"),
        "OXIDATION_FACTOR": ("碳氧化率", "method_parameter", 0.99, "ratio", "authority_published", "省级温室气体清单指南演示因子", "企业演示假设"),
    }
    parameter_sheet = workbook["04_参数与因子"]
    parameter_codes = _codes(workbook, "04_参数与因子")
    for row_index, (code, values) in enumerate(parameters.items(), 3):
        name, role, value, unit, source_mode, source_name, standard_code = values
        payload = {
            "parameter_set_code": parameter_set,
            "parameter_code": code,
            "parameter_name": name,
            "parameter_role": role,
            "value": value,
            "unit": unit,
            "source_mode": source_mode,
            "source_name": source_name,
            "source_version": "2026",
            "standard_code": standard_code,
            "clause_reference": "DEMO-C11",
            "valid_from": "2026-01-01",
            "valid_to": "2026-12-31",
            "reference_year": "2026",
            "geography": "中国",
            "technology": "燃气锅炉",
            "fuel_or_material": "天然气",
            "measurement_method": "演示取值",
            "laboratory": "演示实验室",
            "quality_grade": "A级",
            "uncertainty": "0.02",
            "evidence_code": evidence_code,
            "remarks": "组织碳截图证据自动导入样例",
        }
        for column, field_code in enumerate(parameter_codes, 1):
            parameter_sheet.cell(row=row_index, column=column, value=payload.get(field_code))

    _append(
        workbook,
        "05_证据索引",
        {
            "evidence_code": evidence_code,
            "file_name": f"{evidence_code}.csv",
            "document_type": "csv",
            "document_no": f"DOC-{suffix}",
            "issuer": "演示能源管理系统",
            "period_start": "2026-08-01",
            "period_end": "2026-08-31",
            "storage_key": f".storage/demo-ghg-import/{evidence_code}.csv",
            "sha256": evidence_hash,
            "confidentiality": "internal",
            "verification_status": "verified",
            "valid_from": "2026-08-01",
            "valid_to": "2026-12-31",
        },
    )
    _append(
        workbook,
        "06_数据证据关联",
        {
            "link_code": f"{record_code}-LINK-01",
            "record_type": "activity",
            "record_code": record_code,
            "evidence_code": evidence_code,
            "evidence_role": "activity",
            "page_or_cell_reference": "03_活动数据!A3",
            "coverage_ratio": 1,
        },
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    sys.stdout.write(f"{output}\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build_fixture(args.output)


if __name__ == "__main__":
    main()
