"""Tenant-bound master-data workbook template, validation, and atomic import."""

from __future__ import annotations

# ruff: noqa: EM101, EM102, PERF401, PLR0911, RUF001, TRY003, TRY004
import base64
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import TYPE_CHECKING, Any, TypeVar
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, select

from app.db import models as m
from app.lib.tenancy import TENANT_ID_SESSION_KEY

if TYPE_CHECKING:
    from collections.abc import Iterable

    from openpyxl.worksheet.worksheet import Worksheet
    from sqlalchemy.ext.asyncio import AsyncSession


MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
INTERNAL_KEY_ROW = 2
DATA_START_ROW = 3
T = TypeVar("T")


@dataclass(frozen=True, slots=True)
class FieldSpec:
    key: str
    label: str
    required: bool = False
    field_type: str = "text"
    help_text: str = ""

    @property
    def display_label(self) -> str:
        return f"{self.label}*" if self.required else self.label


SHEET_SPECS: dict[str, tuple[FieldSpec, ...]] = {
    "01_区域": (
        FieldSpec("code", "区域编码", True),
        FieldSpec("name", "区域名称", True),
        FieldSpec("parent_code", "上级区域编码"),
        FieldSpec("area_type", "区域类型"),
        FieldSpec("location_desc", "位置描述"),
        FieldSpec("sort_index", "排序", field_type="integer"),
        FieldSpec("is_enabled", "是否启用", field_type="boolean"),
        FieldSpec("remark", "备注"),
    ),
    "02_设备设备组": (
        FieldSpec("code", "设备编码", True),
        FieldSpec("name", "设备名称", True),
        FieldSpec("asset_kind", "对象类型", True, help_text="meter/device/device_group"),
        FieldSpec("device_type_code", "设备类型编码"),
        FieldSpec("parent_device_code", "上级设备编码"),
        FieldSpec("area_code", "所属区域编码"),
        FieldSpec("external_device_id", "外部设备编号"),
        FieldSpec("asset_code", "资产编号"),
        FieldSpec("warehouse_code", "仓库编码"),
        FieldSpec("warehouse_name", "仓库名称"),
        FieldSpec("energy_type", "能源类型"),
        FieldSpec("protocol", "采集协议"),
        FieldSpec("location", "安装位置"),
        FieldSpec("is_enabled", "是否启用", field_type="boolean"),
        FieldSpec("remark", "备注"),
    ),
    "03_电表绑定": (
        FieldSpec("source_code", "采集源编码", True),
        FieldSpec("topic_filter", "主题", True),
        FieldSpec("gateway_key", "网关键", True),
        FieldSpec("meter_key", "电表键", True),
        FieldSpec("selector_name_pattern", "点位选择规则", True),
        FieldSpec("equipment_code", "电表设备编码", True),
        FieldSpec("external_device_id", "外部设备编号", True),
        FieldSpec("temporary_address", "通信地址", True),
        FieldSpec("identity_status", "身份状态"),
        FieldSpec("scaling_mode", "倍率模式"),
        FieldSpec("multiplier", "倍率", field_type="decimal"),
    ),
    "04_用能系统": (
        FieldSpec("code", "系统编码", True),
        FieldSpec("name", "系统名称", True),
        FieldSpec("system_type", "系统类型", True),
        FieldSpec("parent_code", "上级系统编码"),
        FieldSpec("is_enabled", "是否启用", field_type="boolean"),
        FieldSpec("remark", "备注"),
    ),
    "05_用能系统绑定": (
        FieldSpec("system_code", "系统编码", True),
        FieldSpec("equipment_code", "设备编码"),
        FieldSpec("area_code", "区域编码"),
        FieldSpec("valid_from", "有效起日", field_type="date"),
        FieldSpec("valid_to", "有效止日", field_type="date"),
        FieldSpec("confidence_level", "确认程度"),
        FieldSpec("is_primary", "是否主归属", field_type="boolean"),
        FieldSpec("remark", "备注"),
    ),
    "06_计量边界": (
        FieldSpec("code", "边界编码", True),
        FieldSpec("name", "边界名称", True),
        FieldSpec("boundary_type", "边界类型", True),
        FieldSpec("meter_equipment_code", "电表设备编码"),
        FieldSpec("area_code", "主区域编码"),
        FieldSpec("use_system_code", "用能系统编码"),
        FieldSpec("energy_direction", "计量方向"),
        FieldSpec("include_in_scope2", "纳入范围二", field_type="boolean"),
        FieldSpec("valid_from", "有效起日", field_type="date"),
        FieldSpec("valid_to", "有效止日", field_type="date"),
        FieldSpec("status", "状态"),
        FieldSpec("mapping_version", "映射版本"),
        FieldSpec("source_name", "依据来源"),
        FieldSpec("confidence_level", "确认程度"),
        FieldSpec("remark", "备注"),
    ),
    "07_边界成员": (
        FieldSpec("boundary_code", "边界编码", True),
        FieldSpec("member_code", "成员编码", True),
        FieldSpec("member_type", "成员类型", True, help_text="area/equipment/equipment_group"),
        FieldSpec("area_code", "区域编码"),
        FieldSpec("equipment_code", "设备或设备组编码"),
        FieldSpec("member_role", "成员作用"),
        FieldSpec("allocation_weight", "分配权重", field_type="decimal"),
        FieldSpec("confirmation_status", "确认状态", True, help_text="confirmed/candidate/rejected"),
        FieldSpec("valid_from", "有效起日", field_type="date"),
        FieldSpec("valid_to", "有效止日", field_type="date"),
        FieldSpec("source_name", "依据来源"),
        FieldSpec("remark", "备注"),
    ),
}


@dataclass(slots=True)
class WorkbookData:
    rows: dict[str, list[dict[str, Any]]]
    errors: list[dict[str, Any]]


def master_data_workbook_spec() -> dict[str, Any]:
    return {
        "tenant_source": "current_session",
        "sheets": [
            {
                "name": sheet_name,
                "fields": [
                    {
                        "key": field.key,
                        "label": field.label,
                        "required": field.required,
                        "type": field.field_type,
                        "help_text": field.help_text,
                    }
                    for field in fields
                ],
            }
            for sheet_name, fields in SHEET_SPECS.items()
        ],
    }


def build_master_data_template_bytes() -> bytes:
    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is not None:
        workbook.remove(active_sheet)
    for sheet_name, fields in SHEET_SPECS.items():
        sheet = workbook.create_sheet(sheet_name)
        _format_template_sheet(sheet, fields)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _format_template_sheet(sheet: Worksheet, fields: tuple[FieldSpec, ...]) -> None:
    for column_index, field in enumerate(fields, start=1):
        label_cell = sheet.cell(row=1, column=column_index, value=field.display_label)
        label_cell.font = Font(bold=True, color="FFFFFF")
        label_cell.fill = PatternFill("solid", fgColor="285E61")
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        if field.help_text:
            label_cell.comment = Comment(field.help_text, "nexLCA")
        key_cell = sheet.cell(row=INTERNAL_KEY_ROW, column=column_index, value=field.key)
        key_cell.font = Font(color="777777", italic=True)
        key_cell.fill = PatternFill("solid", fgColor="E8EEF0")
        sheet.column_dimensions[key_cell.column_letter].width = max(14, min(28, len(field.display_label) * 2 + 4))
    sheet.freeze_panes = f"A{DATA_START_ROW}"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(fields)).coordinate}"
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[INTERNAL_KEY_ROW].hidden = True


def _decode_workbook(workbook_base64: str) -> bytes:
    payload = workbook_base64.partition(",")[2] if workbook_base64.startswith("data:") else workbook_base64
    try:
        decoded = base64.b64decode(payload, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("工作簿内容不是有效的 Base64 数据") from exc
    if not decoded:
        raise ValueError("工作簿内容为空")
    if len(decoded) > MAX_WORKBOOK_BYTES:
        raise ValueError("工作簿超过 10 MB 限制")
    return decoded


def parse_master_data_workbook(workbook_base64: str) -> WorkbookData:
    try:
        workbook = load_workbook(BytesIO(_decode_workbook(workbook_base64)), data_only=True, read_only=True)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取工作簿，请使用系统模板") from exc

    rows: dict[str, list[dict[str, Any]]] = {sheet_name: [] for sheet_name in SHEET_SPECS}
    errors: list[dict[str, Any]] = []
    for sheet_name, fields in SHEET_SPECS.items():
        if sheet_name not in workbook.sheetnames:
            errors.append(_error(sheet_name, 0, "sheet", f"缺少工作表 {sheet_name}"))
            continue
        sheet = workbook[sheet_name]
        keys = [str(cell.value or "").strip() for cell in sheet[INTERNAL_KEY_ROW]]
        expected_keys = [field.key for field in fields]
        if keys[: len(expected_keys)] != expected_keys:
            errors.append(_error(sheet_name, INTERNAL_KEY_ROW, "header", "内部列标识与系统模板不一致"))
            continue
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=DATA_START_ROW, values_only=True), start=DATA_START_ROW
        ):
            raw = {field.key: values[index] if index < len(values) else None for index, field in enumerate(fields)}
            if all(value in (None, "") for value in raw.values()):
                continue
            normalized: dict[str, Any] = {"_row": row_number}
            for field in fields:
                try:
                    normalized[field.key] = _coerce_value(raw[field.key], field.field_type)
                except ValueError as exc:
                    errors.append(_error(sheet_name, row_number, field.key, str(exc)))
            for field in fields:
                if field.required and normalized.get(field.key) in (None, ""):
                    errors.append(_error(sheet_name, row_number, field.key, f"{field.label}不能为空"))
            rows[sheet_name].append(normalized)
    return WorkbookData(rows=rows, errors=errors)


def _coerce_value(value: Any, field_type: str) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if field_type == "text":
        return str(value).strip()
    if field_type == "boolean":
        if isinstance(value, bool):
            return value
        normalized = str(value).strip().lower()
        if normalized in {"1", "true", "yes", "y", "是", "启用"}:
            return True
        if normalized in {"0", "false", "no", "n", "否", "停用"}:
            return False
        raise ValueError("应填写是/否")
    if field_type == "integer":
        try:
            return int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError("应填写整数") from exc
    if field_type == "decimal":
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("应填写数值") from exc
    if field_type == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value).strip())
        except ValueError as exc:
            raise ValueError("日期应为 YYYY-MM-DD") from exc
    return value


def _error(sheet: str, row: int, field: str, message: str) -> dict[str, Any]:
    return {"sheet": sheet, "row": row, "field": field, "message": message}


def _tenant_id(session: AsyncSession) -> UUID:
    tenant_id = session.info.get(TENANT_ID_SESSION_KEY)
    if not isinstance(tenant_id, UUID):
        raise ValueError("请先在当前会话中选择租户")
    return tenant_id


async def preview_master_data_workbook(session: AsyncSession, workbook_base64: str) -> dict[str, Any]:
    tenant_id = _tenant_id(session)
    workbook = parse_master_data_workbook(workbook_base64)
    await _validate_references(session, workbook)
    return {
        "tenant_id": str(tenant_id),
        "valid": not workbook.errors,
        "row_count": sum(len(rows) for rows in workbook.rows.values()),
        "sheets": [{"name": sheet_name, "row_count": len(workbook.rows[sheet_name])} for sheet_name in SHEET_SPECS],
        "errors": workbook.errors,
    }


async def _validate_references(session: AsyncSession, workbook: WorkbookData) -> None:  # noqa: PLR0915
    area_rows = workbook.rows["01_区域"]
    device_rows = workbook.rows["02_设备设备组"]
    system_rows = workbook.rows["04_用能系统"]
    boundary_rows = workbook.rows["06_计量边界"]

    await _validate_unique_codes(workbook, "01_区域", area_rows, "code")
    await _validate_unique_codes(workbook, "02_设备设备组", device_rows, "code")
    await _validate_unique_codes(workbook, "04_用能系统", system_rows, "code")
    await _validate_unique_codes(workbook, "06_计量边界", boundary_rows, "code")

    existing_areas = set((await session.scalars(select(m.MonitorSector.code))).all())
    existing_devices = set((await session.scalars(select(m.MonitorEquipment.code))).all())
    existing_systems = set((await session.scalars(select(m.EnergyUseSystem.code))).all())
    existing_boundaries = set((await session.scalars(select(m.MeasurementBoundary.code))).all())
    device_types = set(
        (await session.scalars(select(m.DeviceType.code).where(m.DeviceType.is_enabled.is_(True)))).all()
    )
    source_codes = set((await session.scalars(select(m.MqttDataSource.code))).all())
    subscriptions = set(
        (
            await session.execute(
                select(m.MqttDataSource.code, m.MqttSubscription.topic_filter).join(
                    m.MqttSubscription, m.MqttSubscription.source_id == m.MqttDataSource.id
                )
            )
        ).all()
    )

    area_codes = existing_areas | {row.get("code") for row in area_rows if row.get("code")}
    device_codes = existing_devices | {row.get("code") for row in device_rows if row.get("code")}
    system_codes = existing_systems | {row.get("code") for row in system_rows if row.get("code")}
    boundary_codes = existing_boundaries | {row.get("code") for row in boundary_rows if row.get("code")}

    _validate_tree(workbook, "01_区域", area_rows, "code", "parent_code")
    _validate_tree(workbook, "02_设备设备组", device_rows, "code", "parent_device_code")
    _validate_tree(workbook, "04_用能系统", system_rows, "code", "parent_code")

    for row in area_rows:
        _require_reference(workbook, "01_区域", row, "parent_code", area_codes, "上级区域不存在")
    for row in device_rows:
        _require_reference(workbook, "02_设备设备组", row, "area_code", area_codes, "所属区域不存在")
        _require_reference(workbook, "02_设备设备组", row, "parent_device_code", device_codes, "上级设备或设备组不存在")
        _require_reference(workbook, "02_设备设备组", row, "device_type_code", device_types, "设备类型编码不存在")
        if row.get("asset_kind") not in {"meter", "device", "device_group"}:
            workbook.errors.append(
                _error("02_设备设备组", row["_row"], "asset_kind", "对象类型应为 meter/device/device_group")
            )

    for row in workbook.rows["03_电表绑定"]:
        _require_reference(workbook, "03_电表绑定", row, "source_code", source_codes, "采集源编码不存在")
        if (
            row.get("source_code")
            and row.get("topic_filter")
            and (row["source_code"], row["topic_filter"]) not in subscriptions
        ):
            workbook.errors.append(_error("03_电表绑定", row["_row"], "topic_filter", "当前租户的采集源下不存在该主题"))
        _require_reference(workbook, "03_电表绑定", row, "equipment_code", device_codes, "电表设备编码不存在")

    for row in system_rows:
        _require_reference(workbook, "04_用能系统", row, "parent_code", system_codes, "上级用能系统不存在")
    for row in workbook.rows["05_用能系统绑定"]:
        _require_reference(workbook, "05_用能系统绑定", row, "system_code", system_codes, "用能系统不存在")
        _require_reference(workbook, "05_用能系统绑定", row, "equipment_code", device_codes, "设备不存在")
        _require_reference(workbook, "05_用能系统绑定", row, "area_code", area_codes, "区域不存在")
        if not row.get("equipment_code") and not row.get("area_code"):
            workbook.errors.append(_error("05_用能系统绑定", row["_row"], "equipment_code", "设备和区域至少填写一项"))

    for row in boundary_rows:
        _require_reference(workbook, "06_计量边界", row, "meter_equipment_code", device_codes, "电表设备不存在")
        _require_reference(workbook, "06_计量边界", row, "area_code", area_codes, "主区域不存在")
        _require_reference(workbook, "06_计量边界", row, "use_system_code", system_codes, "用能系统不存在")
        _validate_effective_dates(workbook, "06_计量边界", row)

    for row in workbook.rows["07_边界成员"]:
        _require_reference(workbook, "07_边界成员", row, "boundary_code", boundary_codes, "计量边界不存在")
        _require_reference(workbook, "07_边界成员", row, "area_code", area_codes, "区域不存在")
        _require_reference(workbook, "07_边界成员", row, "equipment_code", device_codes, "设备或设备组不存在")
        expected_field = "area_code" if row.get("member_type") == "area" else "equipment_code"
        if not row.get(expected_field):
            workbook.errors.append(_error("07_边界成员", row["_row"], expected_field, "成员类型与引用对象不一致"))
        if row.get("confirmation_status") not in {"confirmed", "candidate", "rejected"}:
            workbook.errors.append(
                _error("07_边界成员", row["_row"], "confirmation_status", "确认状态应为 confirmed/candidate/rejected")
            )
        _validate_effective_dates(workbook, "07_边界成员", row)


async def _validate_unique_codes(
    workbook: WorkbookData, sheet_name: str, rows: list[dict[str, Any]], field: str
) -> None:
    seen: set[str] = set()
    for row in rows:
        value = row.get(field)
        if value and value in seen:
            workbook.errors.append(_error(sheet_name, row["_row"], field, "工作表内编码重复"))
        elif value:
            seen.add(value)


def _validate_tree(
    workbook: WorkbookData, sheet_name: str, rows: list[dict[str, Any]], code_field: str, parent_field: str
) -> None:
    parent_by_code = {row.get(code_field): row.get(parent_field) for row in rows if row.get(code_field)}
    for row in rows:
        start = row.get(code_field)
        current = start
        visited: set[str] = set()
        while current and current in parent_by_code:
            if current in visited:
                workbook.errors.append(_error(sheet_name, row["_row"], parent_field, "层级关系存在循环"))
                break
            visited.add(current)
            current = parent_by_code.get(current)


def _require_reference(
    workbook: WorkbookData,
    sheet_name: str,
    row: dict[str, Any],
    field: str,
    available: set[Any],
    message: str,
) -> None:
    value = row.get(field)
    if value and value not in available:
        workbook.errors.append(_error(sheet_name, row["_row"], field, message))


def _validate_effective_dates(workbook: WorkbookData, sheet_name: str, row: dict[str, Any]) -> None:
    if row.get("valid_from") and row.get("valid_to") and row["valid_from"] > row["valid_to"]:
        workbook.errors.append(_error(sheet_name, row["_row"], "valid_to", "有效止日不能早于有效起日"))


async def import_master_data_workbook(session: AsyncSession, workbook_base64: str) -> dict[str, Any]:
    tenant_id = _tenant_id(session)
    workbook = parse_master_data_workbook(workbook_base64)
    await _validate_references(session, workbook)
    if workbook.errors:
        raise ValueError(f"工作簿校验失败，共 {len(workbook.errors)} 项错误")

    counters = {sheet_name: {"created": 0, "updated": 0} for sheet_name in SHEET_SPECS}
    try:
        areas = await _upsert_areas(session, tenant_id, workbook.rows["01_区域"], counters["01_区域"])
        devices = await _upsert_devices(
            session, tenant_id, workbook.rows["02_设备设备组"], areas, counters["02_设备设备组"]
        )
        systems = await _upsert_systems(session, tenant_id, workbook.rows["04_用能系统"], counters["04_用能系统"])
        await _upsert_meter_bindings(session, tenant_id, workbook.rows["03_电表绑定"], devices, counters["03_电表绑定"])
        await _upsert_system_bindings(
            session,
            tenant_id,
            workbook.rows["05_用能系统绑定"],
            areas,
            devices,
            counters["05_用能系统绑定"],
        )
        boundaries = await _upsert_boundaries(
            session,
            tenant_id,
            workbook.rows["06_计量边界"],
            areas,
            devices,
            systems,
            counters["06_计量边界"],
        )
        await _upsert_boundary_members(
            session,
            tenant_id,
            workbook.rows["07_边界成员"],
            areas,
            devices,
            boundaries,
            counters["07_边界成员"],
        )
        await session.commit()
    except Exception:
        await session.rollback()
        raise

    return {
        "tenant_id": str(tenant_id),
        "row_count": sum(len(rows) for rows in workbook.rows.values()),
        "sheets": [{"name": name, **counts} for name, counts in counters.items()],
    }


async def _upsert_areas(
    session: AsyncSession, tenant_id: UUID, rows: list[dict[str, Any]], counters: dict[str, int]
) -> dict[str, m.MonitorSector]:
    objects = {obj.code: obj for obj in (await session.scalars(select(m.MonitorSector))).all()}
    for row in _topological_rows(rows, "code", "parent_code"):
        obj = objects.get(row["code"])
        created = obj is None
        if obj is None:
            obj = m.MonitorSector(team_id=tenant_id, code=row["code"])
            session.add(obj)
            objects[obj.code] = obj
        parent = _lookup_code(objects, _optional_code(row, "parent_code"))
        obj.name = row["name"]
        obj.parent = parent
        obj.parent_id = parent.id if parent and parent.id else None
        obj.area_type = row.get("area_type")
        obj.location_desc = row.get("location_desc")
        obj.sort_index = row.get("sort_index") or 0
        obj.is_enable = True if row.get("is_enabled") is None else row["is_enabled"]
        obj.remark = row.get("remark")
        obj.level = (parent.level or 0) + 1 if parent else 0
        obj.path = f"{parent.path}/{obj.code}" if parent and parent.path else f"/{obj.code}"
        await session.flush()
        counters["created" if created else "updated"] += 1
    return objects


async def _upsert_devices(
    session: AsyncSession,
    tenant_id: UUID,
    rows: list[dict[str, Any]],
    areas: dict[str, m.MonitorSector],
    counters: dict[str, int],
) -> dict[str, m.MonitorEquipment]:
    objects = {obj.code: obj for obj in (await session.scalars(select(m.MonitorEquipment))).all()}
    device_types = {obj.code: obj for obj in (await session.scalars(select(m.DeviceType))).all()}
    for row in _topological_rows(rows, "code", "parent_device_code"):
        obj = objects.get(row["code"])
        created = obj is None
        if obj is None:
            obj = m.MonitorEquipment(team_id=tenant_id, code=row["code"])
            session.add(obj)
            objects[obj.code] = obj
        parent = _lookup_code(objects, _optional_code(row, "parent_device_code"))
        area = _lookup_code(areas, _optional_code(row, "area_code"))
        device_type = _lookup_code(device_types, _optional_code(row, "device_type_code"))
        obj.name = row["name"]
        obj.parent_device = parent
        obj.parent_device_id = parent.id if parent and parent.id else None
        obj.sector = area
        obj.sector_id = area.id if area else None
        obj.device_type = device_type
        obj.device_type_id = device_type.id if device_type else None
        obj.category = 0 if row["asset_kind"] == "meter" else None
        for field in (
            "external_device_id",
            "asset_code",
            "warehouse_code",
            "warehouse_name",
            "energy_type",
            "protocol",
            "location",
            "remark",
        ):
            setattr(obj, field, row.get(field))
        obj.is_enable = True if row.get("is_enabled") is None else row["is_enabled"]
        await session.flush()
        counters["created" if created else "updated"] += 1
    return objects


async def _upsert_systems(
    session: AsyncSession, tenant_id: UUID, rows: list[dict[str, Any]], counters: dict[str, int]
) -> dict[str, m.EnergyUseSystem]:
    objects = {obj.code: obj for obj in (await session.scalars(select(m.EnergyUseSystem))).all()}
    for row in _topological_rows(rows, "code", "parent_code"):
        obj = objects.get(row["code"])
        created = obj is None
        if obj is None:
            obj = m.EnergyUseSystem(team_id=tenant_id, code=row["code"])
            session.add(obj)
            objects[obj.code] = obj
        obj.name = row["name"]
        obj.system_type = row["system_type"]
        obj.parent_code = row.get("parent_code")
        obj.is_enabled = True if row.get("is_enabled") is None else row["is_enabled"]
        obj.remark = row.get("remark")
        await session.flush()
        counters["created" if created else "updated"] += 1
    return objects


async def _upsert_meter_bindings(
    session: AsyncSession,
    tenant_id: UUID,
    rows: list[dict[str, Any]],
    devices: dict[str, m.MonitorEquipment],
    counters: dict[str, int],
) -> None:
    sources = {obj.code: obj for obj in (await session.scalars(select(m.MqttDataSource))).all()}
    subscriptions = {
        (source_code, subscription.topic_filter): subscription
        for source_code, subscription in (
            await session.execute(
                select(m.MqttDataSource.code, m.MqttSubscription).join(
                    m.MqttSubscription, m.MqttSubscription.source_id == m.MqttDataSource.id
                )
            )
        ).all()
    }
    del sources
    for row in rows:
        subscription = subscriptions[(row["source_code"], row["topic_filter"])]
        obj = await session.scalar(
            select(m.MqttDeviceBinding).where(
                m.MqttDeviceBinding.subscription_id == subscription.id,
                m.MqttDeviceBinding.gateway_key == row["gateway_key"],
                m.MqttDeviceBinding.meter_key == row["meter_key"],
            )
        )
        created = obj is None
        if obj is None:
            obj = m.MqttDeviceBinding(
                team_id=tenant_id,
                subscription_id=subscription.id,
                gateway_key=row["gateway_key"],
                meter_key=row["meter_key"],
            )
            session.add(obj)
        obj.selector_name_pattern = row["selector_name_pattern"]
        obj.equipment_id = devices[row["equipment_code"]].id
        obj.external_device_id = row["external_device_id"]
        obj.temporary_address = row["temporary_address"]
        obj.identity_status = row.get("identity_status") or "confirmed"
        obj.scaling_mode = row.get("scaling_mode") or "upstream_scaled"
        obj.multiplier = row.get("multiplier") or Decimal(1)
        await session.flush()
        counters["created" if created else "updated"] += 1


async def _upsert_system_bindings(
    session: AsyncSession,
    tenant_id: UUID,
    rows: list[dict[str, Any]],
    areas: dict[str, m.MonitorSector],
    devices: dict[str, m.MonitorEquipment],
    counters: dict[str, int],
) -> None:
    for row in rows:
        equipment = _lookup_code(devices, _optional_code(row, "equipment_code"))
        area = _lookup_code(areas, _optional_code(row, "area_code"))
        conditions = [m.EnergyUseSystemBinding.system_code == row["system_code"]]
        conditions.append(
            m.EnergyUseSystemBinding.equipment_id == equipment.id
            if equipment
            else m.EnergyUseSystemBinding.equipment_id.is_(None)
        )
        conditions.append(
            m.EnergyUseSystemBinding.sector_id == area.id if area else m.EnergyUseSystemBinding.sector_id.is_(None)
        )
        obj = await session.scalar(select(m.EnergyUseSystemBinding).where(and_(*conditions)))
        created = obj is None
        if obj is None:
            obj = m.EnergyUseSystemBinding(team_id=tenant_id, system_code=row["system_code"])
            session.add(obj)
        obj.equipment_id = equipment.id if equipment else None
        obj.equipment_code = equipment.code if equipment else None
        obj.equipment_name = equipment.name if equipment else None
        obj.sector_id = area.id if area else None
        obj.area_id = area.code if area else None
        obj.area_name = area.name if area else None
        obj.valid_from = row.get("valid_from")
        obj.valid_to = row.get("valid_to")
        obj.confidence_level = row.get("confidence_level")
        obj.is_primary = True if row.get("is_primary") is None else row["is_primary"]
        obj.remark = row.get("remark")
        await session.flush()
        counters["created" if created else "updated"] += 1


async def _upsert_boundaries(
    session: AsyncSession,
    tenant_id: UUID,
    rows: list[dict[str, Any]],
    areas: dict[str, m.MonitorSector],
    devices: dict[str, m.MonitorEquipment],
    systems: dict[str, m.EnergyUseSystem],
    counters: dict[str, int],
) -> dict[str, m.MeasurementBoundary]:
    objects = {obj.code: obj for obj in (await session.scalars(select(m.MeasurementBoundary))).all()}
    for row in rows:
        obj = objects.get(row["code"])
        created = obj is None
        if obj is None:
            obj = m.MeasurementBoundary(team_id=tenant_id, code=row["code"], name=row["name"])
            session.add(obj)
            objects[obj.code] = obj
        meter = _lookup_code(devices, _optional_code(row, "meter_equipment_code"))
        area = _lookup_code(areas, _optional_code(row, "area_code"))
        obj.name = row["name"]
        obj.boundary_type = row["boundary_type"]
        obj.meter_equipment_id = meter.id if meter else None
        obj.sector_id = area.id if area else None
        obj.use_system_code = row.get("use_system_code") if row.get("use_system_code") in systems else None
        obj.energy_direction = row.get("energy_direction") or "import"
        obj.include_in_scope2 = True if row.get("include_in_scope2") is None else row["include_in_scope2"]
        obj.valid_from = row.get("valid_from")
        obj.valid_to = row.get("valid_to")
        obj.status = row.get("status") or "active"
        obj.mapping_version = row.get("mapping_version") or "v1"
        obj.source_name = row.get("source_name")
        obj.confidence_level = row.get("confidence_level")
        obj.remark = row.get("remark")
        await session.flush()
        counters["created" if created else "updated"] += 1
    return objects


async def _upsert_boundary_members(
    session: AsyncSession,
    tenant_id: UUID,
    rows: list[dict[str, Any]],
    areas: dict[str, m.MonitorSector],
    devices: dict[str, m.MonitorEquipment],
    boundaries: dict[str, m.MeasurementBoundary],
    counters: dict[str, int],
) -> None:
    for row in rows:
        boundary = boundaries[row["boundary_code"]]
        obj = await session.scalar(
            select(m.MeasurementBoundaryMember).where(
                m.MeasurementBoundaryMember.boundary_id == boundary.id,
                m.MeasurementBoundaryMember.member_code == row["member_code"],
            )
        )
        created = obj is None
        if obj is None:
            obj = m.MeasurementBoundaryMember(
                team_id=tenant_id,
                boundary_id=boundary.id,
                member_code=row["member_code"],
                member_type=row["member_type"],
            )
            session.add(obj)
        area = _lookup_code(areas, _optional_code(row, "area_code"))
        equipment = _lookup_code(devices, _optional_code(row, "equipment_code"))
        obj.member_type = row["member_type"]
        obj.sector_id = area.id if area else None
        obj.equipment_id = equipment.id if equipment else None
        obj.member_role = row.get("member_role") or "covered"
        obj.allocation_weight = row.get("allocation_weight")
        obj.confirmation_status = row["confirmation_status"]
        obj.valid_from = row.get("valid_from")
        obj.valid_to = row.get("valid_to")
        obj.source_name = row.get("source_name")
        obj.remark = row.get("remark")
        await session.flush()
        counters["created" if created else "updated"] += 1


def _topological_rows(rows: list[dict[str, Any]], code_field: str, parent_field: str) -> Iterable[dict[str, Any]]:
    by_code: dict[str, dict[str, Any]] = {str(row[code_field]): row for row in rows}
    visited: set[str] = set()

    def visit(code: str) -> Iterable[dict[str, Any]]:
        if code in visited:
            return []
        row = by_code[code]
        parent_code = _optional_code(row, parent_field)
        ordered: list[dict[str, Any]] = []
        if parent_code in by_code:
            ordered.extend(visit(parent_code))
        visited.add(code)
        ordered.append(row)
        return ordered

    ordered_rows: list[dict[str, Any]] = []
    for code in by_code:
        ordered_rows.extend(visit(code))
    return ordered_rows


def _optional_code(row: dict[str, Any], field: str) -> str | None:
    value = row.get(field)
    return str(value) if value not in (None, "") else None


def _lookup_code(objects: dict[str, T], code: str | None) -> T | None:
    return objects.get(code) if code is not None else None
