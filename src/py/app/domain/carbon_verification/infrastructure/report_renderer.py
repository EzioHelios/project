# ruff: noqa: RUF001, SLF001, TC002, TC003
# pyright: reportUnknownVariableType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportPrivateUsage=false, reportAttributeAccessIssue=false
"""Formal A4 organization-GHG report rendering."""

from __future__ import annotations

import shutil
import subprocess
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.domain.carbon_verification.domain.reporting import ISO_CATEGORY_LABELS, validate_report_context
from app.domain.ghg_inventory.domain.catalog import CATEGORIES

MONTH_LABEL_LENGTH = 7


@dataclass(frozen=True, slots=True)
class RenderedReportFile:
    """One rendered report artifact."""

    path: str
    content: bytes
    content_type: str
    source_name: str


class ReportDocumentRenderer:
    """Render a domestic organization-GHG report from a frozen report context."""

    def __init__(self, *, template_path: Path | None, output_dir: Path) -> None:
        self._template_path = template_path
        self._output_dir = output_dir

    def render(self, report: dict[str, Any]) -> list[RenderedReportFile]:
        validate_report_context(report)
        template = _mapping(report.get("template"))
        if template.get("implementation_status") != "supported" or template.get("renderer_kind") != "generic_a4":
            msg = "所选行业和年度尚未实现正式监管报告适配器"
            raise ValueError(msg)
        self._output_dir.mkdir(parents=True, exist_ok=True)
        year = _safe_filename_part(_mapping(report.get("reporting_period")).get("year") or "all")
        stem = f"{year}-organization-ghg-report"
        docx_path = self._output_dir / f"{stem}.docx"
        pdf_path = self._output_dir / f"{stem}.pdf"
        document = _build_document(report)
        document.save(str(docx_path))
        rendered = [
            RenderedReportFile(
                path=f"report/{docx_path.name}",
                content=docx_path.read_bytes(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                source_name="carbon_verification.organization_report",
            )
        ]
        if _convert_docx_to_pdf(docx_path, pdf_path):
            rendered.append(
                RenderedReportFile(
                    path=f"report/{pdf_path.name}",
                    content=pdf_path.read_bytes(),
                    content_type="application/pdf",
                    source_name="carbon_verification.organization_report",
                )
            )
        return rendered


def _build_document(report: dict[str, Any]) -> DocumentObject:
    document = Document()
    _configure_document(document)
    _cover(document, report)
    document.add_paragraph("目录", style="Title")
    chapters = (
        "一、报告说明与组织概况",
        "二、组织边界与报告边界",
        "三、核算方法、活动数据与参数",
        "四、温室气体排放结果",
        "五、质量管理与不确定性",
        "六、编制与签署",
        "附表一：设备月度活动数据汇总",
        "附表二：资料清单",
    )
    for index, chapter in enumerate(chapters):
        paragraph = document.add_paragraph(chapter + "\t")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), f" PAGEREF chapter{index} \\h ")
        paragraph._p.append(field)
    document.add_page_break()
    _organization_body(document, report, chapters)
    _boundary_body(document, report, chapters)
    _methods_body(document, report, chapters)
    _results_body(document, report, chapters)
    _quality_body(document, report, chapters)
    _appendices_body(document, report, chapters)
    for index, chapter in enumerate(chapters):
        paragraph = next(p for p in document.paragraphs if p.text == chapter and p.style.name == "Heading 1")
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), f"chapter{index}")
        end.set(qn("w:id"), str(index))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)
    return document


def _organization_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    inventory = _mapping(report.get("inventory"))
    enterprise = _mapping(report.get("enterprise"))
    period = _mapping(report.get("reporting_period"))
    details = _mapping(report.get("organization_details"))
    total = _mapping(inventory.get("totals")).get("total_t_co2e")
    _heading(document, chapters[0], 1)
    _prose(
        document,
        f"{enterprise.get('name')}本期温室气体盘查报告的报告期为"
        f"{period.get('period_start')}至{period.get('period_end')}。"
        f"本期纳入边界内的温室气体排放量为{_number(total)} tCO₂e。",
    )
    _narrative(document, report, "purpose", "（一）编制目的与预期用途")
    _narrative(document, report, "intended_users", "（二）预期使用者与披露范围")
    _heading(document, "（三）组织概况", 2)
    _key_value_table(
        document,
        (
            ("报告主体", enterprise.get("name")),
            (
                "组织类型",
                {"enterprise": "企业", "facility": "设施", "group": "集团"}.get(
                    str(enterprise.get("organization_type")), "【待补充：组织类型】"
                ),
            ),
            *(
                (
                    ("统一社会信用代码", enterprise.get("unify_code") or "【待补充】"),
                    ("注册地址", enterprise.get("registered_address") or "【待补充】"),
                    ("法定代表人", details.get("legal_representative") or "【待补充】"),
                )
                if enterprise.get("organization_type") != "facility"
                else ()
            ),
            ("所属单位", enterprise.get("belong_unit")),
            ("行业", enterprise.get("industry") or "【待补充：行业】"),
            ("运营地址", enterprise.get("operation_address")),
            ("联系人", details.get("contact_name") or "【待补充：联系人】"),
            ("联系方式", details.get("contact_details") or "【待补充：联系方式】"),
        ),
    )
    _prose(document, details.get("organization_description") or "【待补充：组织概况】")
    _heading(document, "（四）主要业务及作业过程", 2)
    _prose(document, details.get("process_description") or "【待补充：主要业务及作业过程】")


def _boundary_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    period = _mapping(report.get("reporting_period"))
    source = _mapping(report.get("inventory_source"))
    _heading(document, chapters[1], 1)
    method = {"operational_control": "运营控制法", "financial_control": "财务控制法", "equity_share": "股权比例法"}.get(
        str(source.get("consolidation_method")), "【待补充：组织边界合并方法】"
    )
    _prose(
        document,
        f"本期采用{method}确定组织边界，引用边界为“{source.get('boundary_name')}”。"
        f"报告期覆盖{period.get('period_start')}至{period.get('period_end')}；"
        f"纳入记录的数据覆盖期为{period.get('data_period_start')} 至 {period.get('data_period_end')}。",
    )
    _narrative(document, report, "boundary_rationale", "（一）边界确定与重要性评价")
    _heading(document, "（二）设施及纳入情况", 2)
    _data_table(
        document,
        ("设施 / 区域", "纳入情况", "依据 / 排除理由"),
        [
            (
                row.get("name") or "【待补充：设施名称】",
                "纳入" if row.get("included") else "排除",
                row.get("source") if row.get("included") else row.get("reason"),
            )
            for row in _rows(report, "facilities")
        ],
        widths=(5, 2.5, 8.1),
    )
    _heading(document, "（三）排放源与类别", 2)
    _data_table(
        document,
        ("活动类别", "范围"),
        [
            (_category_label(category), _scope_label(scope))
            for category, scope in sorted(
                {(str(row.get("iso_category_code")), str(row.get("scope"))) for row in _rows(report, "activity_data")}
            )
        ],
        widths=(11, 5),
    )
    _narrative(document, report, "base_year_policy", "（四）基准年与重新计算")
    if period.get("base_year"):
        _prose(document, f"基准年：{period['base_year']}。")


def _methods_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    inventory = _mapping(report.get("inventory"))
    activities = _rows(report, "activity_data")
    _heading(document, chapters[2], 1)
    standard = _mapping(report.get("standard"))
    _prose(
        document,
        f"本次核算采用的规则为{_display(standard.get('name'))}（{_display(standard.get('code'))}）。"
        "排放量按活动数据与选定方法中的因子及参数计算，并以吨二氧化碳当量汇总。",
    )
    _heading(document, "（一）量化方法", 2)
    for index, item in enumerate(_rows(report, "methods"), 1):
        _prose(
            document,
            f"{index}. {_display(item.get('methodName'))}：{item.get('humanFormula') or '【待补充：计算表达式】'}。",
        )
    _prose(document, f"输入能源排放方法：{_scope2_label(inventory.get('default_scope2_method'))}。")
    if inventory.get("gwp_versions"):
        _prose(document, f"GWP版本：{'、'.join(inventory['gwp_versions'])}。")
    _narrative(document, report, "monitoring_procedures", "（二）活动数据收集与监测")
    _data_table(document, ("活动类型", "数量", "单位"), _activity_totals(activities), widths=(8, 4, 3.6))
    _narrative(document, report, "missing_data_treatment", "（三）缺失数据与估算处理")
    _heading(document, "（四）因子及参数", 2)
    factors = _unique_factors(_rows(report, "factor_parameters"))
    _data_table(
        document,
        ("序号", "参数", "数值", "单位", "来源及版本"),
        [
            (
                index,
                row.get("parameter_name") or "【待补充：参数名称】",
                format(_decimal(row["factor_value"]).normalize(), "f")
                if row.get("factor_value") is not None
                else "【待补充：参数数值】",
                row.get("factor_unit"),
                f"{_display(row.get('factor_source'))}\n{_display(row.get('factor_version'))}",
            )
            for index, row in enumerate(factors, 1)
        ],
        widths=(1.2, 2.4, 2.3, 2.5, 7.2),
    )
    for index, factor in enumerate(factors, 1):
        snapshot = _mapping(factor.get("snapshot"))
        codes = {
            row.get("record_code")
            for row in _rows(report, "factor_parameters")
            if _factor_key(row) == _factor_key(factor)
        }
        categories = sorted(
            {_category_label(row.get("iso_category_code")) for row in activities if row.get("record_code") in codes}
        )
        _prose(
            document,
            f"参数{index}：依据{_display(factor.get('standard_code'))}，"
            f"引用位置{_display(factor.get('table_reference') or factor.get('clause_reference'))}；"
            f"适用地区{_display(snapshot.get('geography'))}，参考年份{_display(snapshot.get('referenceYear'))}。"
            f"本期用于{'、'.join(categories) or '【待补充：适用活动类别】'}。",
        )


def _results_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    inventory = _mapping(report.get("inventory"))
    activities = _rows(report, "activity_data")
    total = _mapping(inventory.get("totals")).get("total_t_co2e")
    _heading(document, chapters[3], 1)
    document.paragraphs[-1].paragraph_format.page_break_before = True
    _prose(
        document,
        f"在上述边界与方法下，本报告期温室气体排放总量为{_number(total)} tCO₂e。排放类别、期间变化及计量对象贡献如下。",
    )
    _heading(document, "（一）排放类别汇总", 2)
    groups: dict[tuple[str, str], Decimal] = {}
    for row in activities:
        key = (_category_label(row.get("iso_category_code")), _scope_label(row.get("scope")))
        groups[key] = groups.get(key, Decimal(0)) + _decimal(row.get("computed_co2e")) / 1000
    _data_table(
        document,
        ("活动类别", "范围", "排放量（tCO₂e）"),
        [(category, scope, _number(value)) for (category, scope), value in groups.items()],
        widths=(8, 3.5, 4.5),
    )
    gas_rows = [row for row in _rows(inventory, "gas_summary") if str(row.get("gas_code")).upper() != "CO2E"]
    if gas_rows:
        _summary_table(document, gas_rows, "gas_code", _gas_label)
    _heading(document, "（二）期间排放变化", 2)
    periods: dict[str, Decimal] = {}
    for row in activities:
        label = _period_label(row)
        periods[label] = periods.get(label, Decimal(0)) + _decimal(row.get("computed_co2e")) / Decimal(1000)
    period_rows = sorted(periods.items())
    if all(len(label) == MONTH_LABEL_LENGTH for label, _ in period_rows):
        _caption(document, "图1 月度温室气体排放量")
        _chart(document, period_rows, line=True)
    else:
        _data_table(
            document,
            ("数据期间", "排放量（tCO₂e）"),
            [(label, _number(value)) for label, value in period_rows],
            widths=(11, 5),
        )
    _heading(document, "（三）计量对象排放贡献", 2)
    document.paragraphs[-1].paragraph_format.page_break_before = True
    _caption(document, "图2 计量对象排放贡献及占比")
    _device_contribution(document, activities, _decimal(total))
    _narrative(document, report, "results_commentary", "（四）结果分析")
    if _mapping(report.get("narrative")).get("reduction_actions"):
        _narrative(document, report, "reduction_actions", "（五）减排措施与绩效")


def _quality_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    _heading(document, chapters[4], 1)
    document.paragraphs[-1].paragraph_format.page_break_before = True
    _narrative(document, report, "quality_controls", "（一）质量管理与记录保存")
    _narrative(document, report, "uncertainty_assessment", "（二）不确定性评价")
    activities = _rows(report, "activity_data")
    groups: dict[tuple[str, str], set[str]] = {}
    for row in _rows(report, "quality_issues"):
        key = (str(row.get("title")), str(row.get("action")))
        groups.setdefault(key, set()).add(str(row.get("record_code")))
    pending = sum(row.get("evidence_status") != "ready" for row in activities)
    for title, action in groups:
        if not (pending and title == "核算记录待复核"):
            _prose(document, f"【待补充：{title}；{action}】")
    if pending:
        _prose(document, "【待补充：活动数据、因子参数及原始凭证的复核资料】。")
    _heading(document, chapters[5], 1)
    _prose(document, _mapping(report.get("narrative")).get("prepared_by") or "【待补充：编制负责人及联系方式】")
    _prose(document, "报告主体对本报告所依据资料的真实性、准确性和完整性负责。")
    _signature_grid(document, ("编制人 / 日期", "复核人 / 日期", "批准人 / 日期"))


def _appendices_body(document: DocumentObject, report: dict[str, Any], chapters: tuple[str, ...]) -> None:
    activities = _rows(report, "activity_data")
    _heading(document, chapters[6], 1)
    document.paragraphs[-1].paragraph_format.page_break_before = True
    for (category, unit), rows in _device_month_groups(activities).items():
        _heading(document, f"{_category_label(category)}（{unit}）", 2)
        months = sorted({month for row in rows for month in row["months"]})
        # Six months per panel keep an A4 portrait table readable without reducing type size.
        for offset in range(0, len(months), 6):
            selected = months[offset : offset + 6]
            _data_table(
                document,
                ("设备 / 计量对象", *selected, "小计"),
                [
                    (
                        row["name"],
                        *(_number(row["months"][month]) if month in row["months"] else "—" for month in selected),
                        _number(sum((row["months"].get(month, Decimal(0)) for month in selected), Decimal(0))),
                    )
                    for row in rows
                ]
                + [
                    (
                        "合计",
                        *(
                            _number(sum((row["months"].get(month, Decimal(0)) for row in rows), Decimal(0)))
                            for month in selected
                        ),
                        _number(
                            sum(
                                (row["months"].get(month, Decimal(0)) for row in rows for month in selected), Decimal(0)
                            )
                        ),
                    )
                ],
                widths=(3.7, *(1.7 for _ in selected), 2.1),
                font_size=10.5,
            )
    _heading(document, chapters[7], 1)
    _data_table(
        document,
        ("序号", "文件名称", "资料用途"),
        _evidence_index(_rows(report, "verification_attachments")),
        widths=(1.2, 9.8, 5),
    )


def _prose(document: DocumentObject, text: str) -> None:
    for line in text.splitlines():
        if line.strip():
            paragraph = document.add_paragraph(line.strip())
            paragraph.paragraph_format.first_line_indent = Cm(0.74)


def _gas_label(value: Any) -> str:
    return "CO₂当量（汇总项）" if str(value).upper() == "CO2E" else _uppercase_label(value)


def _narrative(document: DocumentObject, report: dict[str, Any], key: str, title: str) -> None:
    _heading(document, title, 2)
    _prose(document, _mapping(report.get("narrative")).get(key) or f"【待补充：{title.split('）')[-1]}】")


def _factor_key(row: dict[str, Any]) -> tuple[str, ...]:
    return tuple(
        str(row.get(key) or "")
        for key in (
            "parameter_code",
            "factor_value",
            "factor_unit",
            "factor_source",
            "factor_version",
            "standard_code",
            "clause_reference",
            "table_reference",
        )
    )


def _unique_factors(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return list({_factor_key(row): row for row in rows}.values())


def _evidence_index(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    groups: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = (str(row.get("evidence_code")), str(row.get("evidence_role")))
        groups[key] = row
    return [
        (
            index,
            item.get("file_name") or "【待补充：文件名称】",
            "因子来源文件" if item.get("evidence_role") == "factor" else "活动数据凭证",
        )
        for index, item in enumerate(groups.values(), 1)
    ]


def _period_label(row: dict[str, Any]) -> str:
    start, end = str(row.get("period_start") or ""), str(row.get("period_end") or "")
    try:
        begin, finish = date.fromisoformat(start), date.fromisoformat(end)
    except ValueError:
        return "【待补充：数据期间】"
    if (begin.year, begin.month) == (finish.year, finish.month):
        return start[:7]
    return f"{start} 至 {end}"


def _device_month_groups(activities: list[dict[str, Any]]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    groups: dict[tuple[str, str], dict[str, dict[str, Any]]] = {}
    for activity in activities:
        category = (str(activity.get("iso_category_code") or ""), str(activity.get("activity_unit") or ""))
        devices = groups.setdefault(category, {})
        allocations = _rows(activity, "device_allocations")
        if not allocations:
            allocations = [{"device_id": None, "ratio": "1"}]
        ratios = [_decimal(row.get("ratio")) for row in allocations]
        if any(ratio <= 0 for ratio in ratios) or sum(ratios) != Decimal(1):
            msg = "活动数据的设备归集比例不完整，请核实后导出"
            raise ValueError(msg)
        for allocation, ratio in zip(allocations, ratios, strict=True):
            key = str(allocation.get("device_id") or f"unassigned:{activity.get('record_code')}")
            row = devices.setdefault(
                key,
                {
                    "device_id": allocation.get("device_id"),
                    "name": allocation.get("device_name") or "【待补充：设备归属】",
                    "months": {},
                    "co2e_kg": Decimal(0),
                },
            )
            month = _period_label(activity)
            row["months"][month] = (
                row["months"].get(month, Decimal(0)) + _decimal(activity.get("activity_amount")) * ratio
            )
            row["co2e_kg"] += _decimal(activity.get("computed_co2e")) * ratio
    return {
        key: sorted(rows.values(), key=lambda row: (row["device_id"] is None, row["device_id"] or 0))
        for key, rows in groups.items()
    }


def _device_contribution(document: DocumentObject, activities: list[dict[str, Any]], total: Decimal) -> None:
    from app.domain.carbon_verification.infrastructure.report_charts import contribution_bar

    devices: dict[Any, dict[str, Any]] = {}
    for group in _device_month_groups(activities).values():
        for row in group:
            key = row["device_id"] if row["device_id"] is not None else id(row)
            item = devices.setdefault(key, {"name": row["name"], "co2e_kg": Decimal(0)})
            item["co2e_kg"] += row["co2e_kg"]
    rows = list(devices.values())
    rows.sort(key=lambda row: row["co2e_kg"], reverse=True)
    table = _table(document, len(rows) + 1, 4, widths=(6, 3.2, 2.1, 4.7))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, label in enumerate(("计量对象", "排放量\n（tCO₂e）", "占比", "排放贡献")):
        _set_cell(table.cell(0, index), label, header=True)
    maximum = max((abs(row["co2e_kg"]) for row in rows), default=Decimal(0))
    signed = any(row["co2e_kg"] < 0 for row in rows)
    for index, row in enumerate(rows, 1):
        value = row["co2e_kg"] / 1000
        for column, text in enumerate((row["name"], _number(value), _percent(value / total) if total else "—")):
            _set_cell(table.cell(index, column), text)
        table.cell(index, 3).paragraphs[0].add_run().add_picture(
            contribution_bar(row["co2e_kg"], maximum, signed=signed), width=Cm(4)
        )


def _category_label(value: Any) -> str:
    names = {item.code: item.name for item in CATEGORIES}
    return names.get(str(value)) or ISO_CATEGORY_LABELS.get(str(value), f"ISO类别{value}")


def _caption(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    _set_run_font(paragraph.runs[0], "宋体", 10.5)


def _activity_totals(rows: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    groups: dict[tuple[str, str], Decimal] = {}
    for row in rows:
        key = (str(row.get("iso_category_code")), str(row.get("activity_unit")))
        groups[key] = groups.get(key, Decimal(0)) + _decimal(row.get("activity_amount"))
    return [(_category_label(category), _number(amount), unit) for (category, unit), amount in groups.items()]


def _chart(document: DocumentObject, rows: list[tuple[str, Decimal]], *, line: bool = False) -> None:
    from app.domain.carbon_verification.infrastructure.report_charts import result_chart

    if rows:
        document.add_picture(result_chart(rows, line=line), width=Cm(15.6))


def _configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.3)
    section.left_margin, section.right_margin = Cm(2.6), Cm(2.4)
    section.header_distance, section.footer_distance = Cm(1.2), Cm(1.2)
    normal = document.styles["Normal"]
    _style_font(normal, "宋体", 14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(3)
    _style_font(document.styles["Footer"], "宋体", 9)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 14)):
        style = document.styles[name]
        _style_font(style, "宋体", size, bold=True, color="000000")
        _clear_style_paragraph_borders(style)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)
    document.core_properties.title = "组织温室气体排放盘查报告"
    document.core_properties.author = ""
    document.core_properties.comments = ""
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)


def _cover(document: DocumentObject, report: dict[str, Any]) -> None:
    enterprise, period = _mapping(report.get("enterprise")), _mapping(report.get("reporting_period"))
    document.add_paragraph().paragraph_format.space_after = Pt(70)
    subject = document.add_paragraph(str(enterprise.get("name") or "【待补充：报告主体】"), style="Title")
    subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = document.add_paragraph("组织温室气体排放盘查报告", style="Title")
    title.alignment, title.paragraph_format.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(36)
    document.add_paragraph().paragraph_format.space_after = Pt(20)
    rows = (
        ("报告主体（盖章）", enterprise.get("name")),
        ("报告期", f"{_display(period.get('period_start'))} 至 {_display(period.get('period_end'))}"),
        ("报告状态", _report_status(report)),
    )
    _key_value_table(document, rows)
    document.add_page_break()


def _report_status(report: dict[str, Any]) -> str:
    return {
        "draft": "草稿",
        "submitted": "待审核",
        "approved": "已审核",
        "rejected": "已退回",
        "published": "已发布",
        "archived": "已归档",
    }.get(str(_mapping(report.get("inventory_source")).get("status") or ""), "【待确认】")


def _summary_table(document: DocumentObject, rows: list[dict[str, Any]], key: str, labeler: Any) -> None:
    _data_table(
        document,
        ("气体", "排放量（tCO₂e）"),
        [(labeler(row.get(key) or row.get("label")), _number(row.get("t_co2e"))) for row in rows],
        widths=(10, 6),
    )


def _key_value_table(document: DocumentObject, rows: tuple[tuple[str, Any], ...]) -> None:
    table = _table(document, len(rows), 2, widths=(4.2, 11.4))
    for index, (label, value) in enumerate(rows):
        _set_cell(table.cell(index, 0), label, header=True)
        _set_cell(table.cell(index, 1), value)


def _data_table(
    document: DocumentObject,
    headers: tuple[str, ...],
    rows: list[tuple[Any, ...]],
    *,
    widths: tuple[float, ...],
    font_size: float = 10.5,
) -> None:
    table = _table(document, max(1, len(rows)) + 1, len(headers), widths=widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        _set_cell(table.cell(0, index), header, header=True, font_size=font_size)
    if not rows:
        merged = table.cell(1, 0).merge(table.cell(1, len(headers) - 1))
        _set_cell(merged, "无相关记录", font_size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        return
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            _set_cell(table.cell(row_index, column_index), value, font_size=font_size)
    for cell in table.rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True


def _table(document: DocumentObject, rows: int, columns: int, *, widths: tuple[float, ...]) -> Any:
    table = document.add_table(rows=rows, cols=columns)
    table.alignment, table.autofit, table.style = WD_TABLE_ALIGNMENT.CENTER, False, "Table Grid"
    available = 16.0
    widths = tuple(width * available / sum(widths) for width in widths)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, width in enumerate(widths):
            row.cells[index].width = Cm(width)
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(row.cells[index])
    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing = 1
    spacer.paragraph_format.space_after = Pt(4)
    _set_run_font(spacer.add_run(), "宋体", 2)
    spacer.paragraph_format.line_spacing = Pt(2)
    return table


def _set_cell(
    cell: Any, value: Any, *, header: bool = False, font_size: float = 10.5, align: Any = WD_ALIGN_PARAGRAPH.LEFT
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment, paragraph.paragraph_format.space_after = align, Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(_display(value))
    _set_run_font(run, "宋体", font_size, bold=header, color="000000")


def _heading(document: DocumentObject, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _signature_grid(document: DocumentObject, labels: tuple[str, ...]) -> None:
    table = _table(document, 2, len(labels), widths=tuple(5.2 for _ in labels))
    for index, label in enumerate(labels):
        _set_cell(table.cell(0, index), label, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        table.cell(0, index).paragraphs[0].paragraph_format.keep_with_next = True
        _set_cell(table.cell(1, index), "签字：\n日期：", font_size=10.5)
        table.rows[1].height = Cm(2.1)


def _style_font(style: Any, name: str, size: float, *, bold: bool = False, color: str | None = None) -> None:
    style.font.name, style.font.size, style.font.bold = name, Pt(size), bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    _explicit_fonts(style._element.rPr.rFonts)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def _clear_style_paragraph_borders(style: Any) -> None:
    p_pr = style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def _set_run_font(run: Any, name: str, size: float, *, bold: bool = False, color: str | None = None) -> None:
    run.font.name, run.font.size, run.font.bold = name, Pt(size), bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    _explicit_fonts(run._element.rPr.rFonts)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _explicit_fonts(fonts: Any) -> None:
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{key}"), None)
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")


def _set_cell_margins(cell: Any) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if margins.getparent() is None:
        tc_pr.append(margins)
    for name, value in (("top", 80), ("start", 90), ("bottom", 80), ("end", 90)):
        node = margins.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        if node.getparent() is None:
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_page_number(paragraph: Any) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    paragraph._p.append(field)


def _convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> bool:
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if converter is None:
        return False
    result = subprocess.run(  # noqa: S603
        [converter, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
        check=False,
        capture_output=True,
        timeout=90,
    )
    return result.returncode == 0 and pdf_path.exists()


def _mapping(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _rows(data: dict[str, Any], key: str) -> list[dict[str, Any]]:
    value = data.get(key)
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _display(value: Any) -> str:
    if value in (None, ""):
        return "—"
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _number(value: Any) -> str:
    number = _decimal(value)
    if number == number.to_integral():
        return f"{number:,.0f}"
    return f"{number:,.4f}".rstrip("0").rstrip(".")


def _percent(value: Decimal) -> str:
    return f"{(value * Decimal(100)):.1f}%"


def _decimal(value: Any) -> Decimal:
    try:
        return Decimal(str(value or 0))
    except (ArithmeticError, ValueError):
        return Decimal(0)


def _scope_label(value: Any) -> str:
    return {"scope1": "Scope 1", "scope2": "Scope 2", "scope3": "Scope 3"}.get(str(value), _display(value))


def _uppercase_label(value: Any) -> str:
    return _display(value).upper()


def _scope2_label(value: Any) -> str:
    return {
        "location": "位置法",
        "location_based": "位置法",
        "market": "市场法",
        "market_based": "市场法",
        "not_applicable": "不适用",
    }.get(str(value), _display(value))


def _safe_filename_part(value: Any) -> str:
    return "".join(char for char in str(value) if char.isalnum() or char in {"-", "_"}) or "all"
