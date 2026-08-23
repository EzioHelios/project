"""Filtered organization GHG result workbook export."""

from __future__ import annotations

# mypy: disable-error-code=import-untyped
from io import BytesIO
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.domain.ghg_inventory.models import GhgEvidenceFile, GhgEvidenceLink, GhgParameterAssignment
from app.domain.ghg_inventory.queries.analytics import GhgAnalyticsQueryService
from app.domain.ghg_inventory.queries.result_slice import GhgAnalyticsFilter, GhgResultSliceQuery

if TYPE_CHECKING:
    from collections.abc import Iterable

    from openpyxl.worksheet._write_only import WriteOnlyWorksheet
    from sqlalchemy.ext.asyncio import AsyncSession


MAX_EXPORT_ROWS = 50_000


async def build_result_export_bytes(session: AsyncSession, query: GhgAnalyticsFilter) -> bytes:
    result_slice = await GhgResultSliceQuery(session).execute(query)
    if result_slice.total > MAX_EXPORT_ROWS:
        msg = f"当前筛选结果超过{MAX_EXPORT_ROWS}条, 请缩小查询范围后导出"
        raise ValueError(msg)
    analytics = GhgAnalyticsQueryService(session)
    summary = await analytics.results_summary(analytics_filter=query)
    items = await analytics.attributed_result_items(result_slice.rows)
    workbook = Workbook(write_only=True)
    _append_rows(
        workbook.create_sheet("01_查询条件"),
        ("字段", "值"),
        [
            ("核算运行", result_slice.run.run_code),
            ("Scope 2口径", query.scope2_method),
            ("开始日期", query.date_from.isoformat() if query.date_from else ""),
            ("结束日期", query.date_to.isoformat() if query.date_to else ""),
            ("GHG范围", query.ghg_scope or "全部"),
            ("ISO类别", query.iso_category_code or "全部"),
            ("排放源", query.source_code or "全部"),
            ("排放机理", query.source_mechanism or "全部"),
            ("区域ID", query.area_id or "全部"),
            ("设备ID", query.equipment_id or "全部"),
            ("质量状态", query.quality_status or "全部"),
        ],
    )
    _append_rows(
        workbook.create_sheet("02_结果汇总"),
        ("指标", "值", "单位"),
        [
            ("排放合计", summary.get("totalCo2EKg") or "", "kgCO₂e"),
            ("活动数据", summary["activityCount"], "条"),
            ("核算结果", summary["resultCount"], "条"),
            ("缺少证据", summary["missingEvidenceCount"], "条"),
            ("未归集", summary["unallocatedCount"], "条"),
        ],
    )
    _append_rows(
        workbook.create_sheet("03_排放明细"),
        (
            "记录编号",
            "活动名称",
            "开始日期",
            "结束日期",
            "Scope",
            "ISO类别",
            "排放源",
            "排放机理",
            "活动量",
            "活动单位",
            "核算方法",
            "气体",
            "完整排放量(kgCO₂e)",
            "归属比例",
            "归属排放量(kgCO₂e)",
            "区域",
            "设备",
            "质量状态",
            "证据数量",
        ),
        [
            (
                item["recordCode"],
                item["activityName"],
                item["periodStart"],
                item["periodEnd"],
                item["ghgScope"],
                f"{item['isoCategoryCode']} {item['isoCategoryName']}",
                item["sourceName"],
                item["sourceMechanism"],
                item["activityAmount"],
                item["activityUnit"],
                item["methodName"],
                item["gasCode"],
                item["calculatedCo2EKg"],
                item["attributionRatio"],
                item["attributedCo2EKg"],
                "、".join(item["areaNames"]),
                "、".join(item["equipmentNames"]),
                item["qualityStatus"],
                item["evidenceCount"],
            )
            for item in items
        ],
    )
    activity_ids = {row.activity.id for row in result_slice.rows}
    assignments = (
        list(
            (
                await session.scalars(
                    select(GhgParameterAssignment)
                    .where(
                        GhgParameterAssignment.run_id == result_slice.run.id,
                        GhgParameterAssignment.activity_id.in_(activity_ids),
                    )
                    .order_by(GhgParameterAssignment.activity_id, GhgParameterAssignment.parameter_code)
                )
            ).all()
        )
        if activity_ids
        else []
    )
    _append_rows(
        workbook.create_sheet("04_方法与因子"),
        ("活动ID", "参数", "角色", "数值", "单位", "来源", "版本", "标准", "条款"),
        [
            (
                item.activity_id,
                item.parameter_code,
                item.parameter_role,
                str(item.value),
                item.unit,
                item.source_name,
                item.source_version or "",
                item.standard_code or "",
                item.clause_reference or "",
            )
            for item in assignments
        ],
    )
    record_codes = {row.activity.record_code for row in result_slice.rows}
    evidence_rows = (
        list(
            (
                await session.execute(
                    select(GhgEvidenceLink, GhgEvidenceFile)
                    .join(GhgEvidenceFile, GhgEvidenceFile.id == GhgEvidenceLink.evidence_id)
                    .where(
                        GhgEvidenceLink.record_type == "activity",
                        GhgEvidenceLink.record_code.in_(record_codes),
                    )
                    .order_by(GhgEvidenceLink.record_code, GhgEvidenceFile.evidence_code)
                )
            ).all()
        )
        if record_codes
        else []
    )
    _append_rows(
        workbook.create_sheet("05_证据索引"),
        ("记录编号", "证据编号", "文件名", "文档类型", "出具方", "核验状态", "引用位置"),
        [
            (
                link.record_code,
                evidence.evidence_code,
                evidence.file_name,
                evidence.document_type,
                evidence.issuer or "",
                evidence.verification_status,
                link.page_or_cell_reference or "",
            )
            for link, evidence in evidence_rows
        ],
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _append_rows(worksheet: WriteOnlyWorksheet, headers: tuple[str, ...], rows: Iterable[tuple[Any, ...]]) -> None:
    header_cells: list[Any] = []
    for value in headers:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315E6A")
        header_cells.append(cell)
    worksheet.append(header_cells)  # pyright: ignore[reportUnknownMemberType]
    for row in rows:
        worksheet.append(list(row))  # pyright: ignore[reportUnknownMemberType]
