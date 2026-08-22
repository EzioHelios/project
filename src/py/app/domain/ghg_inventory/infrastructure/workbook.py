"""Category-specific XLSX generation and deterministic precheck."""

# mypy: disable-error-code=import-untyped

from __future__ import annotations

# ruff: noqa: RUF001 - Chinese business copy is intentional.
import base64
import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.domain.ghg_inventory.domain.activity_fields import CATEGORY_ACTIVITY_FIELDS
from app.domain.ghg_inventory.domain.allocation import (
    ALLOCATION_ROLES,
    validate_allocation,
    validate_allocation_ratios,
)
from app.domain.ghg_inventory.domain.catalog import CATEGORY_BY_CODE, METHOD_BY_CODE

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

TEMPLATE_VERSION = "GHG-INV-2026.2"
DATA_START_ROW = 3


@dataclass(frozen=True, slots=True)
class ColumnSpec:
    code: str
    label: str
    required: bool = False
    width: int = 18
    choices: tuple[str, ...] = ()


@dataclass(slots=True)
class WorkbookIssue:
    severity: str
    sheet_name: str
    row_number: int | None
    field_code: str | None
    issue_code: str
    message: str
    current_value: str | None = None
    suggestion: str | None = None

    def payload(self) -> dict[str, Any]:
        return {
            "severity": self.severity,
            "sheetName": self.sheet_name,
            "rowNumber": self.row_number,
            "fieldCode": self.field_code,
            "issueCode": self.issue_code,
            "message": self.message,
            "currentValue": self.current_value,
            "suggestion": self.suggestion,
            "status": "open",
        }


@dataclass(slots=True)
class WorkbookPreview:
    boundaries: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    sources: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    activities: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    allocations: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    parameters: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    evidence: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    links: list[dict[str, Any]] = field(default_factory=list[dict[str, Any]])
    issues: list[WorkbookIssue] = field(default_factory=list[WorkbookIssue])
    metadata: dict[str, str] = field(default_factory=dict[str, str])

    def payload(self) -> dict[str, Any]:
        error_count = sum(issue.severity == "error" for issue in self.issues)
        warning_count = sum(issue.severity == "warning" for issue in self.issues)
        global_activity_block = any(issue.severity == "error" and issue.row_number is None for issue in self.issues)
        invalid_activity_rows = {
            issue.row_number
            for issue in self.issues
            if issue.severity == "error" and issue.sheet_name == "03_活动数据" and issue.row_number is not None
        }
        invalid_parameter_rows = {
            issue.row_number
            for issue in self.issues
            if issue.severity == "error" and issue.sheet_name == "04_参数与因子" and issue.row_number is not None
        }
        invalid_parameter_sets = {
            str(row.get("parameter_set_code") or "")
            for row in self.parameters
            if int(row.get("_row_number") or 0) in invalid_parameter_rows
        }
        valid_count = sum(
            not global_activity_block
            and int(row.get("_row_number") or 0) not in invalid_activity_rows
            and str(row.get("parameter_set_code") or "") not in invalid_parameter_sets
            for row in self.activities
        )
        evidence_codes = {str(row.get("evidence_code") or "") for row in self.evidence}
        linked_records = {
            str(row.get("record_code") or "") for row in self.links if row.get("evidence_code") in evidence_codes
        }
        missing_evidence_count = sum(str(row.get("record_code") or "") not in linked_records for row in self.activities)
        return {
            "boundaries": self.boundaries,
            "sources": self.sources,
            "activities": self.activities,
            "allocations": self.allocations,
            "parameters": self.parameters,
            "evidence": self.evidence,
            "links": self.links,
            "issues": [issue.payload() for issue in self.issues],
            "metadata": self.metadata,
            "metrics": {
                "rowCount": len(self.activities),
                "validCount": valid_count,
                "warningCount": warning_count,
                "errorCount": error_count,
                "missingEvidenceCount": missing_evidence_count,
            },
        }


COMMON_ACTIVITY_COLUMNS = (
    ColumnSpec("record_code", "记录编号", True),
    ColumnSpec("source_code", "排放源编号", True),
    ColumnSpec("activity_code", "活动编号", True),
    ColumnSpec("activity_name", "活动名称", True, 24),
    ColumnSpec("period_start", "开始日期", True),
    ColumnSpec("period_end", "结束日期", True),
    ColumnSpec("activity_amount", "活动量", True),
    ColumnSpec("activity_unit", "活动量单位", True),
    ColumnSpec("data_source_type", "数据来源类型", True, choices=("measured", "ledger", "supplier", "estimated")),
    ColumnSpec("measurement_method", "获取/测量方式", True, 24),
    ColumnSpec("parameter_set_code", "参数集编号", True),
    ColumnSpec("national_source_category", "国家标准源类别", False, 24),
    ColumnSpec("boundary_status", "边界状态", True, choices=("included", "excluded", "biogenic_separate", "optional")),
    ColumnSpec("remarks", "备注", False, 24),
)

PACKAGE_FIELDS: dict[str, tuple[ColumnSpec, ...]] = {
    package_code: tuple(
        ColumnSpec(field.code, field.label, field.required, field.width, field.choices) for field in fields
    )
    for package_code, fields in CATEGORY_ACTIVITY_FIELDS.items()
}

PARAMETER_COLUMNS = (
    ColumnSpec("parameter_set_code", "参数集编号", True),
    ColumnSpec("parameter_code", "参数代码", True),
    ColumnSpec("parameter_name", "参数名称", True),
    ColumnSpec("parameter_role", "参数角色", True),
    ColumnSpec("value", "参数值", True),
    ColumnSpec("unit", "单位", True),
    ColumnSpec(
        "source_mode",
        "来源类型",
        True,
        choices=("measured", "supplier", "national_default", "authority_published", "library", "custom_approved"),
    ),
    ColumnSpec("source_name", "来源名称", True, 28),
    ColumnSpec("source_version", "来源版本"),
    ColumnSpec("standard_code", "标准编号"),
    ColumnSpec("clause_reference", "条款"),
    ColumnSpec("table_reference", "表/附录"),
    ColumnSpec("valid_from", "有效期始"),
    ColumnSpec("valid_to", "有效期止"),
    ColumnSpec("reference_year", "基准年份"),
    ColumnSpec("geography", "适用地域"),
    ColumnSpec("technology", "适用技术"),
    ColumnSpec("fuel_or_material", "燃料/物料"),
    ColumnSpec("measurement_method", "检测方法"),
    ColumnSpec("laboratory", "检测机构"),
    ColumnSpec("quality_grade", "质量等级"),
    ColumnSpec("uncertainty", "不确定性"),
    ColumnSpec("evidence_code", "证据编号"),
    ColumnSpec("remarks", "备注"),
)

EVIDENCE_COLUMNS = (
    ColumnSpec("evidence_code", "证据编号", True),
    ColumnSpec("file_name", "文件名称", True, 28),
    ColumnSpec("document_type", "文件类型", True),
    ColumnSpec("document_no", "单据编号"),
    ColumnSpec("issuer", "出具机构", False, 24),
    ColumnSpec("period_start", "覆盖期始"),
    ColumnSpec("period_end", "覆盖期止"),
    ColumnSpec("storage_key", "文件存储标识", True, 28),
    ColumnSpec("sha256", "文件SHA-256", True, 68),
    ColumnSpec("confidentiality", "密级", True, choices=("public", "internal", "confidential")),
    ColumnSpec("verification_status", "核验状态", True, choices=("pending", "verified", "rejected", "expired")),
    ColumnSpec("valid_from", "有效期始"),
    ColumnSpec("valid_to", "有效期止"),
)

LINK_COLUMNS = (
    ColumnSpec("link_code", "关联编号", True),
    ColumnSpec("record_type", "对象类型", True, choices=("activity", "parameter", "source", "boundary")),
    ColumnSpec("record_code", "对象编号", True),
    ColumnSpec("evidence_code", "证据编号", True),
    ColumnSpec(
        "evidence_role",
        "证据角色",
        True,
        choices=("activity", "factor", "boundary", "calibration", "contract", "payment", "quality_control"),
    ),
    ColumnSpec("page_or_cell_reference", "页码/单元格"),
    ColumnSpec("coverage_ratio", "覆盖比例"),
)

ALLOCATION_COLUMNS = (
    ColumnSpec("allocation_code", "归集编号", True),
    ColumnSpec("record_code", "活动记录编号", True),
    ColumnSpec("source_code", "排放源编号", True),
    ColumnSpec("area_code", "区域编码", True),
    ColumnSpec("equipment_code", "设备编码"),
    ColumnSpec("measurement_boundary_code", "计量边界编码", False, 24),
    ColumnSpec("allocation_role", "归集角色", True, 20, tuple(sorted(ALLOCATION_ROLES))),
    ColumnSpec("allocation_ratio", "归集比例", True),
    ColumnSpec("external_facility_code", "外部设施编码"),
    ColumnSpec("external_facility_name", "外部设施名称", False, 24),
    ColumnSpec("source_kind", "归集来源", True, choices=("manual", "source_default", "meter_projection", "migration")),
    ColumnSpec("valid_from", "有效期始"),
    ColumnSpec("valid_to", "有效期止"),
)


def build_template_bytes(context: dict[str, Any]) -> bytes:
    """Build a restrained, category-specific workbook with stable field codes."""
    category = CATEGORY_BY_CODE.get(str(context["iso_category_code"]))
    method = METHOD_BY_CODE.get(str(context["method_code"]))
    if category is None or method is None or method.code not in category.default_methods:
        msg = "分类与计算方法不匹配"
        raise ValueError(msg)

    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is not None:
        workbook.remove(active_sheet)
    instructions = workbook.create_sheet("00_填报说明")
    instructions.append(["组织温室气体活动数据导入模板"])
    instructions.append(["模板版本", TEMPLATE_VERSION])
    instructions.append(["ISO 二级分类", f"{category.code} {category.name}"])
    instructions.append(["计算方法", f"{method.name} / {method.code}"])
    instructions.append(["报告配置", str(context["reporting_profile_code"])])
    instructions.append(["填报规则", "活动数据、参数与证据分别填写；示例不进入正式数据区。"])
    instructions.append(["参数集", "04_参数与因子已列出必填参数；活动数据通过 parameter_set_code 引用同一参数集。"])
    instructions.append(
        ["区域设备归集", "03A_区域设备归集按活动记录填写区域、设备、归集角色和比例；同一活动比例合计必须为100%。"]
    )
    instructions.append(["证据", "05_证据索引登记文件及SHA-256，06_数据证据关联填写对象编号和证据角色。"])
    instructions["A1"].font = Font(size=16, bold=True, color="17324D")
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 72

    metadata = workbook.create_sheet("_metadata")
    for key, value in {
        "template_version": TEMPLATE_VERSION,
        "inventory_code": context["inventory_code"],
        "organization_code": context["organization_code"],
        "facility_code": context["facility_code"],
        "iso_category_code": category.code,
        "reporting_profile_code": context["reporting_profile_code"],
        "method_code": method.code,
        "generated_at": datetime.now().astimezone().isoformat(),
    }.items():
        metadata.append([key, str(value)])
    metadata.sheet_state = "veryHidden"

    boundary_sheet = _add_sheet(
        workbook,
        "01_组织与边界",
        (
            ColumnSpec("inventory_code", "盘查任务", True),
            ColumnSpec("organization_code", "组织编号", True),
            ColumnSpec("facility_code", "设施编号", True),
            ColumnSpec("reporting_profile_code", "报告配置", True),
            ColumnSpec("iso_category_code", "ISO 二级分类", True),
            ColumnSpec("boundary_method", "边界确定方法", True),
        ),
    )
    boundary_codes = [str(cell.value or "") for cell in boundary_sheet[2]]
    boundary_scaffold = {
        "inventory_code": context["inventory_code"],
        "organization_code": context["organization_code"],
        "facility_code": context["facility_code"],
        "reporting_profile_code": context["reporting_profile_code"],
        "iso_category_code": category.code,
    }
    boundary_sheet.append([boundary_scaffold.get(code) for code in boundary_codes])
    _add_sheet(
        workbook,
        "02_排放源",
        (
            ColumnSpec("source_code", "排放源编号", True),
            ColumnSpec("source_name", "排放源名称", True),
            ColumnSpec("activity_type", "活动类型", True),
            ColumnSpec("facility_code", "设施编号", True),
            ColumnSpec("iso_category_code", "ISO 二级分类", True),
        ),
    )
    _add_sheet(workbook, "03_活动数据", COMMON_ACTIVITY_COLUMNS + PACKAGE_FIELDS[category.package_code])
    _add_sheet(workbook, "03A_区域设备归集", _allocation_columns(context))
    parameter_sheet = _add_sheet(workbook, "04_参数与因子", PARAMETER_COLUMNS)
    parameter_codes = [column.code for column in PARAMETER_COLUMNS]
    for parameter in method.parameters:
        if not parameter["required"]:
            continue
        scaffold = {
            "parameter_set_code": "PARAMSET-01",
            "parameter_code": parameter["code"],
            "parameter_name": parameter["name"],
            "parameter_role": parameter["role"],
            "unit": parameter["units"][0],
        }
        parameter_sheet.append([scaffold.get(code) for code in parameter_codes])
    _add_sheet(workbook, "05_证据索引", EVIDENCE_COLUMNS)
    _add_sheet(workbook, "06_数据证据关联", LINK_COLUMNS)
    national = _add_sheet(
        workbook,
        "07_国家格式对照",
        (
            ColumnSpec("reporting_profile_id", "报告配置"),
            ColumnSpec("standard_code", "标准编号"),
            ColumnSpec("canonical_entity", "规范化对象"),
            ColumnSpec("canonical_field", "规范化字段"),
            ColumnSpec("required_level", "要求级别"),
            ColumnSpec("formula_method_version", "方法版本"),
        ),
    )
    national.append(
        [
            context["reporting_profile_code"],
            "GB/T 32150-2025",
            "ghg_activity_record",
            "iso_category_code",
            "required",
            f"{method.code}/2026.1",
        ]
    )
    _add_sheet(
        workbook,
        "08_校验结果",
        (
            ColumnSpec("severity", "级别"),
            ColumnSpec("sheet_name", "工作表"),
            ColumnSpec("row_number", "行号"),
            ColumnSpec("field_code", "字段"),
            ColumnSpec("issue_code", "错误代码"),
            ColumnSpec("message", "问题"),
            ColumnSpec("suggestion", "修复建议", width=36),
        ),
    )
    dictionary = _add_sheet(
        workbook,
        "09_枚举字典",
        (ColumnSpec("dictionary", "字典"), ColumnSpec("code", "代码"), ColumnSpec("name", "名称")),
    )
    for category_item in CATEGORY_BY_CODE.values():
        dictionary.append(["iso_category", category_item.code, category_item.name])
    for method_item in METHOD_BY_CODE.values():
        dictionary.append(["method", method_item.code, method_item.name])
    for sheet_name in ("07_国家格式对照", "08_校验结果", "09_枚举字典"):
        workbook[sheet_name].sheet_state = "hidden"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _allocation_columns(context: dict[str, Any]) -> tuple[ColumnSpec, ...]:
    dynamic_choices = {
        "area_code": tuple(str(value) for value in context.get("area_codes", []) if value),
        "equipment_code": tuple(str(value) for value in context.get("equipment_codes", []) if value),
        "measurement_boundary_code": tuple(
            str(value) for value in context.get("measurement_boundary_codes", []) if value
        ),
    }
    return tuple(
        ColumnSpec(
            column.code,
            column.label,
            column.required,
            column.width,
            dynamic_choices.get(column.code, column.choices),
        )
        for column in ALLOCATION_COLUMNS
    )


def _add_sheet(workbook: Workbook, name: str, columns: tuple[ColumnSpec, ...]) -> Worksheet:
    sheet = workbook.create_sheet(name)
    sheet.append([column.label + ("*" if column.required else "") for column in columns])
    sheet.append([column.code for column in columns])
    sheet.row_dimensions[2].hidden = True
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}2"
    for index, column in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="17324D")
        cell.fill = PatternFill("solid", fgColor="E8F0F3" if not column.required else "D7E9EB")
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = column.width
        if column.choices:
            validation = DataValidation(type="list", formula1='"' + ",".join(column.choices) + '"')
            sheet.add_data_validation(validation)
            validation.add(  # pyright: ignore[reportUnknownMemberType] - openpyxl stubs omit CellRange typing
                f"{get_column_letter(index)}3:{get_column_letter(index)}1002"
            )
    return sheet


def decode_and_preview(workbook_base64: str, expected_context: dict[str, Any]) -> tuple[WorkbookPreview, str]:
    """Decode, hash, validate, and parse a workbook without executing formulas."""
    try:
        raw = base64.b64decode(workbook_base64, validate=True)
        workbook = load_workbook(BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        msg = "无法读取工作簿，请上传系统生成的 .xlsx 模板"
        raise ValueError(msg) from exc
    preview = WorkbookPreview()
    preview.metadata = _read_metadata(workbook)
    _validate_metadata(preview, expected_context)
    preview.boundaries = _read_rows(workbook, "01_组织与边界", preview.issues)
    preview.sources = _read_rows(workbook, "02_排放源", preview.issues)
    preview.activities = _read_rows(workbook, "03_活动数据", preview.issues)
    preview.allocations = _read_rows(workbook, "03A_区域设备归集", preview.issues)
    preview.parameters = _read_rows(workbook, "04_参数与因子", preview.issues)
    preview.evidence = _read_rows(workbook, "05_证据索引", preview.issues)
    preview.links = _read_rows(workbook, "06_数据证据关联", preview.issues)
    _validate_boundary_and_sources(preview, expected_context)
    _validate_activities(preview, expected_context)
    _validate_allocations(preview, expected_context)
    _validate_parameters(preview, expected_context)
    _validate_evidence(preview)
    return preview, hashlib.sha256(raw).hexdigest()


def _read_metadata(workbook: Workbook) -> dict[str, str]:
    if "_metadata" not in workbook.sheetnames:
        return {}
    return {
        str(row[0].value): str(row[1].value or "")
        for row in workbook["_metadata"].iter_rows(min_row=1, max_col=2)
        if row[0].value
    }


def _validate_metadata(preview: WorkbookPreview, expected: dict[str, Any]) -> None:
    for key in (
        "template_version",
        "inventory_code",
        "organization_code",
        "facility_code",
        "iso_category_code",
        "reporting_profile_code",
        "method_code",
    ):
        expected_value = TEMPLATE_VERSION if key == "template_version" else str(expected.get(key) or "")
        actual = preview.metadata.get(key)
        if actual != expected_value:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "_metadata",
                    None,
                    key,
                    "TEMPLATE_CONTEXT_MISMATCH",
                    f"模板中的{key}与当前导入上下文不一致",
                    actual,
                    "重新下载当前上下文的模板",
                )
            )


def _read_rows(workbook: Workbook, sheet_name: str, issues: list[WorkbookIssue]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        issues.append(WorkbookIssue("error", sheet_name, None, None, "MISSING_SHEET", f"缺少工作表 {sheet_name}"))
        return []
    sheet = workbook[sheet_name]
    codes = [str(cell.value or "").strip() for cell in sheet[2]]
    rows: list[dict[str, Any]] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=DATA_START_ROW), DATA_START_ROW):
        values = [cell.value for cell in cells[: len(codes)]]
        if not any(value not in (None, "") for value in values):
            continue
        row = {code: _json_value(value) for code, value in zip(codes, values, strict=False) if code}
        row["_row_number"] = row_number
        rows.append(row)
    return rows


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _required(row: dict[str, Any], fields: tuple[str, ...], sheet: str, issues: list[WorkbookIssue]) -> None:
    issues.extend(
        WorkbookIssue(
            "error",
            sheet,
            int(row["_row_number"]),
            field_code,
            "REQUIRED_FIELD_MISSING",
            "缺少必填字段",
            None,
            "补充后重新预检",
        )
        for field_code in fields
        if row.get(field_code) in (None, "")
    )


def _validate_activities(preview: WorkbookPreview, context: dict[str, Any]) -> None:
    category = CATEGORY_BY_CODE[str(context["iso_category_code"])]
    if not preview.activities:
        preview.issues.append(
            WorkbookIssue("error", "03_活动数据", None, None, "NO_ACTIVITY_ROWS", "没有可导入的活动数据")
        )
        return
    seen: set[str] = set()
    parameter_sets = {str(row.get("parameter_set_code")) for row in preview.parameters if row.get("parameter_set_code")}
    for row in preview.activities:
        _required(
            row,
            tuple(
                column.code
                for column in COMMON_ACTIVITY_COLUMNS + PACKAGE_FIELDS[category.package_code]
                if column.required
            ),
            "03_活动数据",
            preview.issues,
        )
        record_code = str(row.get("record_code") or "")
        if record_code in seen:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03_活动数据",
                    int(row["_row_number"]),
                    "record_code",
                    "DUPLICATE_RECORD_CODE",
                    "记录编号在工作簿内重复",
                    record_code,
                )
            )
        seen.add(record_code)
        try:
            amount = Decimal(str(row.get("activity_amount")))
        except (InvalidOperation, TypeError):
            amount = Decimal(-1)
        if amount < 0:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03_活动数据",
                    int(row["_row_number"]),
                    "activity_amount",
                    "INVALID_ACTIVITY_AMOUNT",
                    "活动量必须为非负数",
                    str(row.get("activity_amount")),
                )
            )
        invalid_period = False
        try:
            period_start = date.fromisoformat(str(row.get("period_start")))
            period_end = date.fromisoformat(str(row.get("period_end")))
        except (TypeError, ValueError):
            invalid_period = True
        else:
            invalid_period = period_start > period_end
        if invalid_period:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03_活动数据",
                    int(row["_row_number"]),
                    "period_start",
                    "INVALID_ACTIVITY_PERIOD",
                    "活动开始日期和结束日期必须有效，且开始日期不得晚于结束日期",
                )
            )
        parameter_set = str(row.get("parameter_set_code") or "")
        if parameter_set not in parameter_sets:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03_活动数据",
                    int(row["_row_number"]),
                    "parameter_set_code",
                    "PARAMETER_SET_NOT_FOUND",
                    "未找到活动记录引用的参数集",
                    parameter_set,
                )
            )
        row.update(
            {
                "inventory_code": context["inventory_code"],
                "organization_code": context["organization_code"],
                "facility_code": context["facility_code"],
                "iso_category_code": category.code,
                "ghg_scope": category.ghg_scope,
                "scope3_category": category.scope3_category,
                "source_mechanism": category.source_mechanism,
                "value_chain_position": category.value_chain_position,
            }
        )


def _validate_boundary_and_sources(preview: WorkbookPreview, context: dict[str, Any]) -> None:
    if len(preview.boundaries) != 1:
        preview.issues.append(
            WorkbookIssue(
                "error",
                "01_组织与边界",
                None,
                None,
                "BOUNDARY_ROW_REQUIRED",
                "组织与边界工作表必须且只能填写一行",
            )
        )
    for row in preview.boundaries:
        _required(
            row,
            (
                "inventory_code",
                "organization_code",
                "facility_code",
                "reporting_profile_code",
                "iso_category_code",
                "boundary_method",
            ),
            "01_组织与边界",
            preview.issues,
        )
        for field_code in (
            "inventory_code",
            "organization_code",
            "facility_code",
            "reporting_profile_code",
            "iso_category_code",
        ):
            if str(row.get(field_code) or "") != str(context.get(field_code) or ""):
                preview.issues.append(
                    WorkbookIssue(
                        "error",
                        "01_组织与边界",
                        int(row["_row_number"]),
                        field_code,
                        "BOUNDARY_CONTEXT_MISMATCH",
                        "组织与边界字段与当前导入上下文不一致",
                        str(row.get(field_code) or ""),
                    )
                )
    if not preview.sources:
        preview.issues.append(
            WorkbookIssue("error", "02_排放源", None, None, "SOURCE_ROW_REQUIRED", "至少登记一个排放源")
        )
    seen_sources: set[str] = set()
    for row in preview.sources:
        _required(
            row,
            ("source_code", "source_name", "activity_type", "facility_code", "iso_category_code"),
            "02_排放源",
            preview.issues,
        )
        source_code = str(row.get("source_code") or "")
        if source_code in seen_sources:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "02_排放源",
                    int(row["_row_number"]),
                    "source_code",
                    "DUPLICATE_SOURCE_CODE",
                    "排放源编号重复",
                    source_code,
                )
            )
        seen_sources.add(source_code)
        for field_code in ("facility_code", "iso_category_code"):
            if str(row.get(field_code) or "") != str(context.get(field_code) or ""):
                preview.issues.append(
                    WorkbookIssue(
                        "error",
                        "02_排放源",
                        int(row["_row_number"]),
                        field_code,
                        "SOURCE_CONTEXT_MISMATCH",
                        "排放源字段与当前导入上下文不一致",
                        str(row.get(field_code) or ""),
                    )
                )
    for row in preview.activities:
        if str(row.get("source_code") or "") not in seen_sources:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03_活动数据",
                    int(row["_row_number"]),
                    "source_code",
                    "SOURCE_NOT_FOUND",
                    "活动记录引用的排放源不存在",
                    str(row.get("source_code") or ""),
                )
            )


def _validate_allocations(preview: WorkbookPreview, context: dict[str, Any]) -> None:
    category_code = str(context["iso_category_code"])
    activity_sources = {
        str(row.get("record_code") or ""): str(row.get("source_code") or "") for row in preview.activities
    }
    if not preview.allocations:
        preview.issues.append(
            WorkbookIssue(
                "error",
                "03A_区域设备归集",
                None,
                None,
                "ALLOCATION_REQUIRED",
                "每条活动数据必须完成区域归集",
            )
        )
        return
    grouped: dict[str, list[Decimal]] = {}
    seen_codes: set[str] = set()
    for row in preview.allocations:
        _required(
            row,
            (
                "allocation_code",
                "record_code",
                "source_code",
                "area_code",
                "allocation_role",
                "allocation_ratio",
                "source_kind",
            ),
            "03A_区域设备归集",
            preview.issues,
        )
        allocation_code = str(row.get("allocation_code") or "")
        record_code = str(row.get("record_code") or "")
        source_code = str(row.get("source_code") or "")
        if allocation_code in seen_codes:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03A_区域设备归集",
                    int(row["_row_number"]),
                    "allocation_code",
                    "DUPLICATE_ALLOCATION_CODE",
                    "归集编号重复",
                    allocation_code,
                )
            )
        seen_codes.add(allocation_code)
        if record_code not in activity_sources:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03A_区域设备归集",
                    int(row["_row_number"]),
                    "record_code",
                    "ALLOCATION_ACTIVITY_NOT_FOUND",
                    "归集记录引用的活动不存在",
                    record_code,
                )
            )
        elif source_code != activity_sources[record_code]:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03A_区域设备归集",
                    int(row["_row_number"]),
                    "source_code",
                    "ALLOCATION_SOURCE_MISMATCH",
                    "归集记录的排放源与活动数据不一致",
                    source_code,
                )
            )
        try:
            ratio = Decimal(str(row.get("allocation_ratio")))
        except (InvalidOperation, TypeError):
            ratio = Decimal(0)
        grouped.setdefault(record_code, []).append(ratio)
        try:
            validate_allocation(
                category_code=category_code,
                role=str(row.get("allocation_role") or ""),
                area_present=bool(str(row.get("area_code") or "").strip()),
                equipment_present=bool(str(row.get("equipment_code") or "").strip()),
                external_present=bool(str(row.get("external_facility_code") or "").strip()),
            )
        except ValueError as exc:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03A_区域设备归集",
                    int(row["_row_number"]),
                    "allocation_role",
                    "INVALID_ALLOCATION_POLICY",
                    str(exc),
                )
            )
    for record_code in activity_sources:
        try:
            validate_allocation_ratios(grouped.get(record_code, []))
        except ValueError as exc:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "03A_区域设备归集",
                    None,
                    "allocation_ratio",
                    "INVALID_ALLOCATION_TOTAL",
                    f"活动 {record_code}：{exc}",
                )
            )


def _validate_parameters(preview: WorkbookPreview, context: dict[str, Any]) -> None:
    method = METHOD_BY_CODE[str(context["method_code"])]
    required_codes = {str(item["code"]) for item in method.parameters if item["required"]}
    definitions = {str(item["code"]): item for item in method.parameters}
    grouped: dict[str, set[str]] = {}
    seen: set[tuple[str, str]] = set()
    for row in preview.parameters:
        _required(
            row,
            (
                "parameter_set_code",
                "parameter_code",
                "parameter_name",
                "parameter_role",
                "value",
                "unit",
                "source_mode",
                "source_name",
            ),
            "04_参数与因子",
            preview.issues,
        )
        code = str(row.get("parameter_code") or "")
        parameter_set = str(row.get("parameter_set_code") or "")
        key = (parameter_set, code)
        if key in seen:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "04_参数与因子",
                    int(row["_row_number"]),
                    "parameter_code",
                    "DUPLICATE_PARAMETER",
                    "同一参数集中不能重复填写同一参数",
                    code,
                )
            )
        seen.add(key)
        grouped.setdefault(parameter_set, set()).add(code)
        definition = definitions.get(code)
        if code and definition is None:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "04_参数与因子",
                    int(row["_row_number"]),
                    "parameter_code",
                    "PARAMETER_NOT_ALLOWED",
                    f"参数 {code} 不属于所选计算方法 {method.code}",
                    code,
                )
            )
        if definition is not None:
            allowed_units = {str(value) for value in definition["units"]}
            if str(row.get("unit") or "") not in allowed_units:
                preview.issues.append(
                    WorkbookIssue(
                        "error",
                        "04_参数与因子",
                        int(row["_row_number"]),
                        "unit",
                        "PARAMETER_UNIT_NOT_ALLOWED",
                        f"参数 {code} 的单位必须为：{', '.join(sorted(allowed_units))}",
                        str(row.get("unit") or ""),
                    )
                )
            allowed_sources = {str(value) for value in definition["sourceModes"]}
            if str(row.get("source_mode") or "") not in allowed_sources:
                preview.issues.append(
                    WorkbookIssue(
                        "error",
                        "04_参数与因子",
                        int(row["_row_number"]),
                        "source_mode",
                        "PARAMETER_SOURCE_NOT_ALLOWED",
                        f"参数 {code} 的来源类型不受当前方法支持",
                        str(row.get("source_mode") or ""),
                    )
                )
        try:
            Decimal(str(row.get("value")))
        except (InvalidOperation, TypeError):
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "04_参数与因子",
                    int(row["_row_number"]),
                    "value",
                    "INVALID_PARAMETER_VALUE",
                    "参数值不是有效数值",
                    str(row.get("value")),
                )
            )
        if row.get("source_mode") == "measured" and (
            not row.get("measurement_method") or not row.get("laboratory") or not row.get("evidence_code")
        ):
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "04_参数与因子",
                    int(row["_row_number"]),
                    "source_mode",
                    "MEASURED_PARAMETER_EVIDENCE_REQUIRED",
                    "实测参数必须填写检测方法、机构和证据编号",
                )
            )
    for parameter_set in {str(row.get("parameter_set_code") or "") for row in preview.activities}:
        for missing_code in sorted(required_codes - grouped.get(parameter_set, set())):
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "04_参数与因子",
                    None,
                    "parameter_code",
                    "REQUIRED_PARAMETER_MISSING",
                    f"参数集 {parameter_set} 缺少必需参数 {missing_code}",
                )
            )
    if not method.executable:
        preview.issues.append(
            WorkbookIssue(
                "error",
                "04_参数与因子",
                None,
                "method_code",
                "METHOD_NOT_EXECUTABLE",
                "当前方法尚未发布为可执行计算方法",
                method.code,
                "选择已发布的可执行方法",
            )
        )


def _validate_evidence(preview: WorkbookPreview) -> None:
    evidence_codes: set[str] = set()
    for row in preview.evidence:
        _required(
            row,
            (
                "evidence_code",
                "file_name",
                "document_type",
                "storage_key",
                "sha256",
                "confidentiality",
                "verification_status",
            ),
            "05_证据索引",
            preview.issues,
        )
        code = str(row.get("evidence_code") or "")
        if code in evidence_codes:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "05_证据索引",
                    int(row["_row_number"]),
                    "evidence_code",
                    "DUPLICATE_EVIDENCE_CODE",
                    "证据编号重复",
                    code,
                )
            )
        evidence_codes.add(code)
        digest = str(row.get("sha256") or "")
        if len(digest) != hashlib.sha256().digest_size * 2 or any(
            char not in "0123456789abcdefABCDEF" for char in digest
        ):
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "05_证据索引",
                    int(row["_row_number"]),
                    "sha256",
                    "INVALID_SHA256",
                    "文件哈希必须为64位SHA-256十六进制值",
                    digest,
                )
            )
    record_codes = {str(row.get("record_code") or "") for row in preview.activities}
    parameter_sets = {str(row.get("parameter_set_code") or "") for row in preview.parameters}
    source_codes = {str(row.get("source_code") or "") for row in preview.sources}
    boundary_codes = {str(row.get("inventory_code") or "") for row in preview.boundaries}
    valid_targets = {
        "activity": (record_codes, "ACTIVITY_NOT_FOUND", "关联的活动记录不存在"),
        "parameter": (parameter_sets, "PARAMETER_SET_NOT_FOUND", "关联的参数集不存在"),
        "source": (source_codes, "SOURCE_NOT_FOUND", "关联的排放源不存在"),
        "boundary": (boundary_codes, "BOUNDARY_NOT_FOUND", "组织边界证据的对象编号必须为当前盘查任务编号"),
    }
    linked_activities: set[str] = set()
    link_codes: set[str] = set()
    for row in preview.links:
        _required(
            row,
            ("link_code", "record_type", "record_code", "evidence_code", "evidence_role"),
            "06_数据证据关联",
            preview.issues,
        )
        record_type = str(row.get("record_type") or "")
        record_code = str(row.get("record_code") or "")
        evidence_code = str(row.get("evidence_code") or "")
        link_code = str(row.get("link_code") or "")
        _validate_evidence_target(row, record_type, record_code, valid_targets, preview.issues)
        if link_code in link_codes:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "06_数据证据关联",
                    int(row["_row_number"]),
                    "link_code",
                    "DUPLICATE_EVIDENCE_LINK_CODE",
                    "关联编号重复",
                    link_code,
                )
            )
        link_codes.add(link_code)
        if evidence_code not in evidence_codes:
            preview.issues.append(
                WorkbookIssue(
                    "error",
                    "06_数据证据关联",
                    int(row["_row_number"]),
                    "evidence_code",
                    "EVIDENCE_NOT_FOUND",
                    "关联的证据编号不存在",
                    evidence_code,
                )
            )
        if record_type == "activity" and evidence_code in evidence_codes:
            linked_activities.add(record_code)
        if row.get("coverage_ratio") not in (None, ""):
            try:
                coverage = Decimal(str(row["coverage_ratio"]))
            except (InvalidOperation, TypeError):
                coverage = Decimal(-1)
            if coverage < 0 or coverage > 1:
                preview.issues.append(
                    WorkbookIssue(
                        "error",
                        "06_数据证据关联",
                        int(row["_row_number"]),
                        "coverage_ratio",
                        "INVALID_COVERAGE_RATIO",
                        "证据覆盖比例必须在0至1之间",
                        str(row.get("coverage_ratio")),
                    )
                )
    for row in preview.activities:
        record_code = str(row.get("record_code") or "")
        if record_code not in linked_activities:
            preview.issues.append(
                WorkbookIssue(
                    "warning",
                    "06_数据证据关联",
                    int(row["_row_number"]),
                    "record_code",
                    "ACTIVITY_EVIDENCE_MISSING",
                    "活动记录尚未关联证据",
                    record_code,
                    "在证据索引中登记文件并建立活动关联",
                )
            )


def _validate_evidence_target(
    row: dict[str, Any],
    record_type: str,
    record_code: str,
    valid_targets: dict[str, tuple[set[str], str, str]],
    issues: list[WorkbookIssue],
) -> None:
    target = valid_targets.get(record_type)
    if target is None:
        issues.append(
            WorkbookIssue(
                "error",
                "06_数据证据关联",
                int(row["_row_number"]),
                "record_type",
                "EVIDENCE_RECORD_TYPE_NOT_ALLOWED",
                "导入证据只能关联活动、参数集、排放源或组织边界",
                record_type,
            )
        )
        return
    valid_codes, issue_code, message = target
    if record_code not in valid_codes:
        issues.append(
            WorkbookIssue(
                "error",
                "06_数据证据关联",
                int(row["_row_number"]),
                "record_code",
                issue_code,
                message,
                record_code,
            )
        )
