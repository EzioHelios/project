from __future__ import annotations

import base64
import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.domain.carbon.real_company_package_importer import (
    DATASET_PATHS,
    RealCompanyPackageImportResult,
    import_real_company_datasets,
)

if TYPE_CHECKING:
    from collections.abc import Mapping

    from openpyxl.worksheet.worksheet import Worksheet
    from sqlalchemy.ext.asyncio import AsyncSession


SANITIZED_ENTERPRISE_CODE = "honda_sanitized"
SANITIZED_TENANT_SLUG = "sample-honda-metal-2025"
SANITIZED_ENTERPRISE_NAME = "本田金属样例企业(脱敏)"
SANITIZED_SOURCE_NAME = "本田金属样例资料包(脱敏)"
INTERNAL_KEY_ROW = 2
DATA_START_ROW = 3


@dataclass(frozen=True, slots=True)
class WorkbookFieldSpec:
    """User-facing field metadata for one import column."""

    key: str
    label: str
    field_type: str = "text"
    required: bool = False
    example: str = ""
    help_text: str = ""
    choices: dict[str, str] | None = None

    @property
    def display_label(self) -> str:
        return f"{self.label}*" if self.required else self.label

    def to_payload(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "label": self.label,
            "type": self.field_type,
            "required": self.required,
            "example": self.example,
            "help_text": self.help_text,
            "choices": self.choices or {},
        }

WORKBOOK_SHEETS: dict[str, str] = {
    "01_企业基本信息": "enterprise_profiles",
    "02_报告组织": "reporting_organizations",
    "03_报告期间": "reporting_periods",
    "04_组织边界": "organization_boundaries",
    "05_产品主数据": "production_products",
    "06_产量分母": "production_output_facts",
    "07_能源计量事实": "energy_interval_facts",
    "08_Scope1移动源燃烧": "scope1_mobile_combustion",
    "09_Scope1固定源燃烧": "scope1_stationary_combustion",
    "10_Scope1制冷剂逸散": "scope1_refrigerant_leak",
    "11_Scope2电力账单": "scope2_electricity_bills",
    "12_Scope2外购热汽冷": "scope2_purchased_energy",
    "13_Scope3废弃物": "scope3_waste_disposal",
    "14_Scope3外购运输": "scope3_third_party_transport",
    "15_排放因子": "emission_factors",
    "16_PCF研究": "product_footprint_studies",
    "17_PCF系统边界": "product_system_boundaries",
    "18_PCF生命周期阶段": "product_lifecycle_stages",
    "19_PCF清单项": "product_footprint_inventory_items",
    "20_PCF结果": "product_footprint_results",
    "21_组织碳核算批次": "carbon_accounting_batches",
    "22_组织碳清册结果": "carbon_emission_inventory_items",
}

DATASET_HEADERS: dict[str, list[str]] = {
    "enterprise_profiles": [
        "enterprise_code",
        "enterprise_name",
        "tenant_slug",
        "source_package",
        "package_root",
        "reporting_year",
        "data_role",
        "standardization_status",
        "platform_tables",
    ],
    "reporting_organizations": [
        "enterprise_code",
        "code",
        "name",
        "organization_type",
        "control_method",
        "standard_code",
        "source_name",
        "target_table",
    ],
    "reporting_periods": [
        "enterprise_code",
        "period_code",
        "reporting_org_code",
        "name",
        "period_type",
        "period_start",
        "period_end",
        "status",
        "source_name",
        "target_table",
    ],
    "organization_boundaries": [
        "enterprise_code",
        "boundary_no",
        "reporting_org_code",
        "period_code",
        "boundary_name",
        "boundary_type",
        "boundary_method",
        "operational_boundary",
        "status",
        "standard_code",
        "source_name",
        "target_table",
    ],
    "production_products": [
        "enterprise_code",
        "product_code",
        "product_name",
        "product_spec",
        "category",
        "output_kind",
        "default_unit",
        "standard_unit",
        "remark",
        "target_table",
    ],
    "production_output_facts": [
        "enterprise_code",
        "source_record_id",
        "product_code",
        "product_name",
        "period_start",
        "period_end",
        "good_qty",
        "unit",
        "source_system",
        "record_status",
        "remark",
        "target_table",
    ],
    "energy_interval_facts": [
        "enterprise_code",
        "source_record_id",
        "energy_type_code",
        "equipment_code",
        "equipment_name",
        "area_name",
        "use_system_code",
        "use_system_name",
        "granularity",
        "period_start",
        "period_end",
        "amount",
        "unit",
        "data_quality",
        "source_system",
        "remark",
        "target_table",
    ],
    "scope1_mobile_combustion": [
        "enterprise_code",
        "record_date",
        "asset_id",
        "asset_name",
        "fuel_type",
        "amount",
        "unit",
        "mileage_km",
        "work_hours",
        "attachment",
        "remark",
        "target_table",
    ],
    "scope1_stationary_combustion": [
        "enterprise_code",
        "facility_name",
        "period_start",
        "period_end",
        "fuel_type",
        "amount",
        "unit",
        "attachment",
        "remark",
        "target_table",
    ],
    "scope1_refrigerant_leak": [
        "enterprise_code",
        "device_id",
        "device_name",
        "fill_date",
        "refrigerant_type",
        "amount_kg",
        "leak_type",
        "attachment",
        "remark",
        "target_table",
    ],
    "scope2_electricity_bills": [
        "enterprise_code",
        "account_number",
        "billing_month",
        "total_kwh",
        "amount_cny",
        "attachment",
        "remark",
        "target_table",
    ],
    "scope2_purchased_energy": [
        "enterprise_code",
        "energy_type_code",
        "energy_type_name",
        "supplier_name",
        "period_start",
        "period_end",
        "amount",
        "unit",
        "standard_amount",
        "standard_unit",
        "tce_amount",
        "amount_cny",
        "area_id",
        "area_name",
        "facility_name",
        "attachment",
        "remark",
        "calc_status",
        "factor_source",
        "factor_year",
        "factor_unit",
        "activity_amount",
        "activity_unit",
        "computed_co2e",
        "computation_note",
        "standard_code",
        "conversion_factor_source",
        "target_table",
    ],
    "scope3_waste_disposal": [
        "enterprise_code",
        "date",
        "category",
        "weight_ton",
        "treatment_method",
        "vendor_name",
        "attachment",
        "remark",
        "target_table",
    ],
    "scope3_third_party_transport": [
        "enterprise_code",
        "record_date",
        "waybill_no",
        "origin",
        "destination",
        "vehicle_type",
        "distance_km",
        "load_ton",
        "fuel_amount",
        "attachment",
        "remark",
        "target_table",
    ],
    "emission_factors": [
        "enterprise_code",
        "scope",
        "category",
        "sub_category",
        "unit",
        "factor_value",
        "factor_unit",
        "year",
        "source",
        "source_system",
        "source_snapshot_id",
        "projection_status",
        "review_status",
        "is_active",
        "note",
        "target_table",
    ],
    "product_footprint_studies": [
        "enterprise_code",
        "study_code",
        "study_name",
        "product_code",
        "product_name",
        "product_spec",
        "standard_code",
        "standard_name",
        "functional_unit",
        "declared_unit",
        "period_start",
        "period_end",
        "boundary_type",
        "status",
        "source_name",
        "confidence_level",
        "remark",
        "target_table",
    ],
    "product_system_boundaries": [
        "enterprise_code",
        "study_code",
        "boundary_type",
        "included_stages",
        "excluded_stages",
        "geography",
        "technology_scope",
        "temporal_scope",
        "standard_code",
        "source_name",
        "target_table",
    ],
    "product_lifecycle_stages": [
        "enterprise_code",
        "study_code",
        "stage_code",
        "stage_name",
        "sequence",
        "included",
        "description",
        "stage_co2e",
        "stage_share",
        "target_table",
    ],
    "product_footprint_inventory_items": [
        "enterprise_code",
        "study_code",
        "item_no",
        "stage_code",
        "activity_type",
        "item_name",
        "activity_amount",
        "activity_unit",
        "factor_value",
        "factor_unit",
        "factor_source",
        "allocation_ratio",
        "computed_co2e",
        "calc_status",
        "data_quality",
        "source_type",
        "source_ref",
        "evidence_path",
        "included",
        "remark",
        "target_table",
    ],
    "product_footprint_results": [
        "enterprise_code",
        "study_code",
        "result_version",
        "total_co2e",
        "unit",
        "footprint_per_declared_unit",
        "result_status",
        "inventory_count",
        "missing_factor_count",
        "stage_payload",
        "result_payload",
        "report_status",
        "remark",
        "target_table",
    ],
    "carbon_accounting_batches": [
        "enterprise_code",
        "batch_no",
        "batch_name",
        "reporting_org_code",
        "period_code",
        "boundary_no",
        "status",
        "inventory_status",
        "standard_code",
        "source_name",
        "total_co2e",
        "record_count",
        "included_record_count",
        "remark",
        "target_table",
    ],
    "carbon_emission_inventory_items": [
        "enterprise_code",
        "batch_no",
        "record_type",
        "record_key",
        "scope",
        "record_label",
        "display_name",
        "period",
        "source_kind",
        "activity_amount",
        "activity_unit",
        "computed_co2e",
        "factor_source",
        "calc_status",
        "quality_status",
        "included",
        "source_payload",
        "target_table",
    ],
}

FIELD_LABELS: dict[str, str] = {
    "account_number": "电力户号",
    "activity_amount": "活动数据量",
    "activity_type": "活动类型",
    "activity_unit": "活动数据单位",
    "address": "地址",
    "allocation_ratio": "分摊比例",
    "amount": "用量",
    "amount_cny": "金额(元)",
    "amount_kg": "数量(kg)",
    "area_id": "区域编码",
    "area_name": "区域名称",
    "asset_id": "资产编码",
    "asset_name": "资产名称",
    "attachment": "附件路径",
    "batch_name": "核算批次名称",
    "batch_no": "核算批次号",
    "billing_month": "账期",
    "boundary_method": "边界确定方法",
    "boundary_name": "边界名称",
    "boundary_no": "边界编号",
    "boundary_type": "边界类型",
    "calc_status": "计算状态",
    "category": "类别",
    "code": "组织编码",
    "computed_co2e": "排放量(kgCO2e)",
    "computation_note": "计算说明",
    "confidence_level": "置信等级",
    "control_method": "控制方法",
    "conversion_factor_source": "折标来源",
    "data_quality": "数据质量",
    "data_role": "数据角色",
    "data_source": "数据来源",
    "date": "日期",
    "declared_unit": "声明单位",
    "default_unit": "默认单位",
    "description": "说明",
    "destination": "终点",
    "device_id": "设备编码",
    "device_name": "设备名称",
    "display_name": "展示名称",
    "distance_km": "运输距离(km)",
    "energy_type_code": "能源品种",
    "energy_type_name": "能源品种名称",
    "enterprise_code": "企业编码",
    "enterprise_name": "企业名称",
    "equipment_code": "设备编码",
    "equipment_name": "设备名称",
    "evidence_path": "证据路径",
    "excluded_stages": "排除阶段",
    "factor_source": "因子来源",
    "factor_unit": "因子单位",
    "factor_value": "因子值",
    "factor_year": "因子年份",
    "facility_name": "设施/系统名称",
    "file_name": "文件名",
    "fill_date": "填充/泄漏日期",
    "footprint_per_declared_unit": "单位产品足迹",
    "fuel_amount": "燃料用量",
    "fuel_type": "燃料类型",
    "functional_unit": "功能单位",
    "geography": "地理范围",
    "good_qty": "合格产量",
    "granularity": "时间粒度",
    "included": "是否纳入",
    "included_record_count": "纳入记录数",
    "included_stages": "纳入阶段",
    "is_active": "是否启用",
    "inventory_count": "清单项数量",
    "inventory_status": "清册状态",
    "item_name": "清单项名称",
    "item_no": "清单项编号",
    "leak_type": "类型",
    "load_ton": "载重(吨)",
    "material_code": "物料编码",
    "material_name": "物料名称",
    "mileage_km": "里程(km)",
    "missing_factor_count": "缺因子数量",
    "name": "名称",
    "note": "备注",
    "operational_boundary": "运营边界",
    "organization_type": "组织类型",
    "origin": "起点",
    "output_kind": "产出类型",
    "package_root": "资料包路径",
    "period": "期间",
    "period_code": "期间编码",
    "period_end": "结束日期",
    "period_start": "开始日期",
    "period_type": "期间类型",
    "platform_tables": "平台表",
    "projection_status": "投影状态",
    "product_code": "产品编码",
    "product_name": "产品名称",
    "product_spec": "规格型号",
    "record_count": "记录数",
    "record_date": "记录日期",
    "record_key": "记录键",
    "record_label": "记录标签",
    "record_status": "记录状态",
    "record_type": "记录类型",
    "quality_status": "质量状态",
    "refrigerant_type": "制冷剂类型",
    "remark": "备注",
    "report_status": "报告状态",
    "reporting_org_code": "报告组织编码",
    "reporting_year": "报告年度",
    "result_payload": "结果明细",
    "result_status": "结果状态",
    "result_version": "结果版本",
    "review_status": "审核状态",
    "scope": "范围",
    "sequence": "顺序",
    "source": "来源",
    "source_name": "来源名称",
    "source_kind": "来源类型",
    "source_package": "来源资料包",
    "source_payload": "来源明细",
    "source_record_id": "来源记录号",
    "source_ref": "来源引用",
    "source_snapshot_id": "来源快照号",
    "source_system": "来源系统",
    "source_type": "来源类型",
    "stage_code": "阶段编码",
    "stage_co2e": "阶段排放量",
    "stage_name": "生命周期阶段",
    "stage_payload": "阶段明细",
    "stage_share": "阶段占比",
    "standard_amount": "折标量",
    "standard_code": "标准编号",
    "standard_name": "标准名称",
    "standard_unit": "标准单位",
    "standardization_status": "标准化状态",
    "status": "状态",
    "study_code": "研究编码",
    "study_name": "研究名称",
    "sub_category": "子类别",
    "supplier_name": "供应商",
    "target_table": "目标表",
    "technology_scope": "技术范围",
    "temporal_scope": "时间范围",
    "tce_amount": "折标煤(tce)",
    "tenant_slug": "租户标识",
    "total_co2e": "总排放量",
    "total_kwh": "总电量(kWh)",
    "treatment_method": "处置方式",
    "unit": "单位",
    "use_system_code": "用能系统编码",
    "use_system_name": "用能系统名称",
    "vehicle_type": "车辆类型",
    "vendor_name": "处置单位",
    "waybill_no": "运单号",
    "weight_ton": "重量(吨)",
    "work_hours": "作业小时",
    "year": "年份",
}

FIELD_TYPES: dict[str, str] = {
    "amount": "number",
    "amount_cny": "number",
    "amount_kg": "number",
    "activity_amount": "number",
    "computed_co2e": "number",
    "distance_km": "number",
    "factor_value": "number",
    "factor_year": "number",
    "fill_date": "date",
    "good_qty": "number",
    "included_record_count": "number",
    "load_ton": "number",
    "mileage_km": "number",
    "period_end": "date",
    "period_start": "date",
    "record_date": "date",
    "total_kwh": "number",
    "weight_ton": "number",
    "work_hours": "number",
    "year": "number",
}

DATASET_REQUIRED_FIELDS: dict[str, set[str]] = {
    "enterprise_profiles": {"enterprise_code", "enterprise_name", "tenant_slug"},
    "reporting_organizations": {"code", "name"},
    "reporting_periods": {"period_code", "reporting_org_code", "period_start", "period_end"},
    "organization_boundaries": {"boundary_no", "reporting_org_code", "period_code", "boundary_name"},
    "production_products": {"product_code", "product_name"},
    "production_output_facts": {"source_record_id", "product_code", "period_start", "period_end", "good_qty", "unit"},
    "energy_interval_facts": {"source_record_id", "energy_type_code", "period_start", "period_end", "amount", "unit"},
    "scope1_mobile_combustion": {"record_date", "fuel_type", "amount", "unit"},
    "scope1_stationary_combustion": {"facility_name", "period_start", "period_end", "fuel_type", "amount", "unit"},
    "scope1_refrigerant_leak": {"fill_date", "refrigerant_type", "amount_kg", "leak_type"},
    "scope2_electricity_bills": {"account_number", "billing_month", "total_kwh"},
    "scope2_purchased_energy": {"energy_type_code", "period_start", "period_end", "amount", "unit"},
    "scope3_waste_disposal": {"date", "category", "weight_ton", "treatment_method"},
    "scope3_third_party_transport": {"waybill_no", "origin", "destination", "vehicle_type", "distance_km", "load_ton"},
    "emission_factors": {"scope", "category", "sub_category", "unit", "factor_value", "year"},
    "product_footprint_studies": {"study_code", "study_name", "product_code"},
    "product_lifecycle_stages": {"study_code", "stage_code", "stage_name"},
    "product_footprint_inventory_items": {"study_code", "item_no", "stage_code", "item_name"},
    "product_footprint_results": {"study_code", "result_version", "total_co2e"},
    "carbon_accounting_batches": {"batch_no", "batch_name", "reporting_org_code", "period_code"},
    "carbon_emission_inventory_items": {"batch_no", "record_type", "record_key", "scope"},
}

FIELD_CHOICES: dict[str, dict[str, str]] = {
    "energy_type_code": {"电力": "electricity", "天然气": "natural_gas", "外购热力": "purchased_heat", "外购蒸汽": "purchased_steam", "外购冷量": "purchased_cooling"},
    "fuel_type": {"柴油": "diesel", "汽油": "gasoline", "液化天然气(LNG)": "lng", "压缩天然气(CNG)": "cng", "天然气": "天然气"},
    "leak_type": {"填充": "fill", "泄漏": "leak", "维修补充": "repair"},
    "treatment_method": {"回收": "recycle", "能源回收": "energy_recovery", "填埋": "landfill", "焚烧": "incineration"},
    "vehicle_type": {"重型货车": "heavy_truck", "中型货车": "medium_truck", "轻型货车": "light_truck", "厢式货车": "van"},
    "granularity": {"小时": "hour", "日": "day", "月": "month"},
    "data_quality": {"实测": "measured", "账单": "bill", "估算": "estimated"},
    "record_status": {"已确认": "confirmed", "草稿": "draft"},
    "status": {"草稿": "draft", "已通过": "approved", "已发布": "published", "有效": "active", "完成": "completed"},
    "included": {"是": "true", "否": "false"},
    "is_active": {"是": "true", "否": "false"},
}

FIELD_EXAMPLES: dict[str, str] = {
    "account_number": "电力户号001",
    "amount": "86.5",
    "amount_kg": "12.5",
    "billing_month": "2025-01",
    "energy_type_code": "电力",
    "fill_date": "2025-02-20",
    "fuel_type": "柴油",
    "period_end": "2025-01-31",
    "period_start": "2025-01-01",
    "record_date": "2025-01-15",
    "total_kwh": "12000",
    "unit": "L/kWh/t",
}

DATASET_SAMPLE_ROWS: dict[str, list[dict[str, str]]] = {
    "scope1_mobile_combustion": [
        {
            "record_date": "2025-01-15",
            "asset_id": "MOBILE-SAMPLE-01",
            "asset_name": "脱敏柴油叉车",
            "fuel_type": "diesel",
            "amount": "86.5000",
            "unit": "L",
            "mileage_km": "",
            "work_hours": "42.50",
            "attachment": "",
            "remark": "脱敏补充样例, 用于验证移动源燃烧导入链路, 不代表本田源数据。",
        }
    ],
    "scope1_refrigerant_leak": [
        {
            "device_id": "REF-SAMPLE-01",
            "device_name": "脱敏冷库机组",
            "fill_date": "2025-02-20",
            "refrigerant_type": "R410A",
            "amount_kg": "12.5000",
            "leak_type": "repair",
            "attachment": "",
            "remark": "脱敏补充样例, 用于验证制冷剂逸散导入链路, 不代表本田源数据。",
        }
    ],
    "scope3_third_party_transport": [
        {
            "record_date": "2025-03-10",
            "waybill_no": "TP-SAMPLE-20250310-001",
            "origin": "脱敏供应商仓",
            "destination": "脱敏厂区",
            "vehicle_type": "heavy_truck",
            "distance_km": "168.50",
            "load_ton": "8.2000",
            "fuel_amount": "",
            "attachment": "",
            "remark": "脱敏补充样例, 用于验证第三方运输导入链路, 不代表本田源数据。",
        }
    ],
}


async def import_data_collection_workbook(
    session: AsyncSession,
    workbook_bytes: bytes,
    *,
    dry_run: bool = False,
    package_root: str | Path = ".",
) -> RealCompanyPackageImportResult:
    """Import one system-shaped multi-sheet energy-carbon data collection workbook."""

    datasets = parse_data_collection_workbook(workbook_bytes)
    return await import_real_company_datasets(session, datasets, package_root=package_root, dry_run=dry_run)


async def import_data_collection_workbook_base64(
    session: AsyncSession,
    workbook_base64: str,
    *,
    dry_run: bool = False,
    package_root: str | Path = ".",
) -> RealCompanyPackageImportResult:
    """Decode and import a base64-encoded workbook payload."""

    try:
        workbook_bytes = base64.b64decode(workbook_base64, validate=True)
    except ValueError as exc:
        msg = "workbook_base64 不是合法的 base64 内容"
        raise ValueError(msg) from exc
    return await import_data_collection_workbook(session, workbook_bytes, dry_run=dry_run, package_root=package_root)


def preview_data_collection_workbook_base64(workbook_base64: str) -> dict[str, Any]:
    """Decode and preview a base64-encoded workbook payload."""

    try:
        workbook_bytes = base64.b64decode(workbook_base64, validate=True)
    except ValueError as exc:
        msg = "workbook_base64 不是合法的 base64 内容"
        raise ValueError(msg) from exc
    return preview_data_collection_workbook(workbook_bytes)


def build_data_collection_template_bytes() -> bytes:
    """Build a user-facing blank workbook template with sample rows for supplemental tables."""

    workbook = _build_collection_workbook({})
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def data_collection_workbook_spec() -> dict[str, Any]:
    """Return sheet and field metadata for frontend import workflows."""

    return {
        "sheets": [
            {
                "sheet_name": sheet_name,
                "dataset_name": dataset_name,
                "fields": [field.to_payload() for field in _field_specs_for_dataset(dataset_name, DATASET_SAMPLE_ROWS.get(dataset_name, []))],
            }
            for sheet_name, dataset_name in WORKBOOK_SHEETS.items()
        ]
    }


def preview_data_collection_workbook(workbook_bytes: bytes) -> dict[str, Any]:
    """Preview workbook coverage, row counts, and first rows without importing."""

    datasets = parse_data_collection_workbook(workbook_bytes)
    sheets = []
    for sheet_name, dataset_name in WORKBOOK_SHEETS.items():
        rows = datasets.get(dataset_name, [])
        fields = _field_specs_for_dataset(dataset_name, rows)
        missing_required = _missing_required_fields(dataset_name, rows)
        sheets.append(
            {
                "sheet_name": sheet_name,
                "dataset_name": dataset_name,
                "row_count": len(rows),
                "fields": [field.to_payload() for field in fields],
                "missing_required": missing_required,
                "preview_rows": [_display_row(dataset_name, row) for row in rows[:5]],
            }
        )
    return {"sheets": sheets, "total_rows": sum(item["row_count"] for item in sheets)}


def parse_data_collection_workbook(workbook_bytes: bytes) -> dict[str, list[dict[str, str]]]:
    """Read known workbook sheets into the normalized dataset contract used by package imports."""

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    datasets = {dataset_name: [] for dataset_name in DATASET_PATHS}
    for sheet_name, dataset_name in WORKBOOK_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        datasets[dataset_name] = _worksheet_rows(workbook[sheet_name])
    if not datasets["enterprise_profiles"]:
        msg = "工作簿缺少 01_企业基本信息 或该 sheet 无数据"
        raise ValueError(msg)
    return datasets


def build_collection_workbook_from_package(package_dir: str | Path, output_path: str | Path) -> Path:
    """Build a sanitized workbook from an existing standardized real-company CSV package."""

    package_root = Path(package_dir).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    datasets: dict[str, list[dict[str, str]]] = {}
    for dataset_name in WORKBOOK_SHEETS.values():
        rows = _read_package_dataset(package_root, dataset_name)
        if not rows:
            rows = DATASET_SAMPLE_ROWS.get(dataset_name, [])
        datasets[dataset_name] = [_sanitize_row(dataset_name, row) for row in rows]

    workbook = _build_collection_workbook(datasets)
    workbook.save(output)
    return output


def _worksheet_rows(sheet: Worksheet) -> list[dict[str, str]]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(rows) if any(_cell_to_text(cell) for cell in row)), None)
    if header_index is None:
        return []
    sheet_name = sheet.title
    dataset_name = WORKBOOK_SHEETS.get(sheet_name, "")
    if len(rows) > 1 and _looks_like_internal_key_row(rows[INTERNAL_KEY_ROW - 1], dataset_name):
        headers = [_cell_to_text(cell).lstrip("\ufeff").strip() for cell in rows[INTERNAL_KEY_ROW - 1]]
        data_start_index = DATA_START_ROW - 1
    else:
        raw_headers = [_cell_to_text(cell).lstrip("\ufeff").strip() for cell in rows[header_index]]
        headers = [_resolve_field_key(dataset_name, header) for header in raw_headers]
        data_start_index = header_index + 1
    data_rows: list[dict[str, str]] = []
    for row in rows[data_start_index:]:
        values = [_cell_to_text(cell) for cell in row[: len(headers)]]
        if not any(values):
            continue
        normalized_row: dict[str, str] = {}
        for header, value in zip(headers, values, strict=False):
            if not header:
                continue
            normalized_row[header] = _internal_value(header, value)
        data_rows.append(normalized_row)
    return data_rows


def _build_collection_workbook(datasets: dict[str, list[dict[str, str]]]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_metadata_sheet(workbook)
    for sheet_name, dataset_name in WORKBOOK_SHEETS.items():
        rows = datasets.get(dataset_name, DATASET_SAMPLE_ROWS.get(dataset_name, []))
        _append_dataset_sheet(workbook, sheet_name, dataset_name, rows)
    return workbook


def _append_metadata_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("00_填报说明")
    rows = [
        ("字段", "内容"),
        ("模板用途", "能碳数据采集包, 按系统数据表拆分为多 sheet"),
        ("样例来源", SANITIZED_SOURCE_NAME),
        ("脱敏口径", "企业/租户标识已替换, 本地附件路径已清空, 数值仅用于导入联调"),
        ("导入范围", "企业基础信息, 报告范围, 产量, 能源, 组织碳活动数据, 排放因子, 产品碳足迹"),
    ]
    for row in rows:
        sheet.append(row)
    _format_sheet(sheet)


def _append_dataset_sheet(workbook: Workbook, sheet_name: str, dataset_name: str, rows: list[dict[str, str]]) -> None:
    sheet = workbook.create_sheet(sheet_name)
    fields = _field_specs_for_dataset(dataset_name, rows)
    headers = [field.key for field in fields]
    sheet.append([field.display_label for field in fields])
    sheet.append(headers)
    sheet.row_dimensions[INTERNAL_KEY_ROW].hidden = True
    for column_index, field in enumerate(fields, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7" if field.required else "F6F8FA")
        comment_lines = [f"类型: {field.field_type}"]
        if field.example:
            comment_lines.append(f"示例: {field.example}")
        if field.choices:
            comment_lines.append("可选值: " + ", ".join(field.choices))
        if field.help_text:
            comment_lines.append(field.help_text)
        cell.comment = Comment("\n".join(comment_lines), "系统")
    for row in rows:
        sheet.append([_display_value(header, row.get(header, "")) for header in headers])
    _format_sheet(sheet)


def _headers_for_dataset(dataset_name: str, rows: list[dict[str, str]]) -> list[str]:
    headers = list(DATASET_HEADERS.get(dataset_name, []))
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    return headers


def _field_specs_for_dataset(dataset_name: str, rows: list[dict[str, str]]) -> list[WorkbookFieldSpec]:
    required = DATASET_REQUIRED_FIELDS.get(dataset_name, set())
    return [
        WorkbookFieldSpec(
            key=key,
            label=FIELD_LABELS.get(key, _fallback_label(key)),
            field_type=FIELD_TYPES.get(key, "text"),
            required=key in required,
            example=FIELD_EXAMPLES.get(key, ""),
            choices=FIELD_CHOICES.get(key),
        )
        for key in _headers_for_dataset(dataset_name, rows)
    ]


def _format_sheet(sheet: Worksheet) -> None:
    sheet.freeze_panes = f"A{DATA_START_ROW}"
    if sheet.max_row and sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = min(max(len(_cell_to_text(cell.value)) for cell in column_cells), 36)
        sheet.column_dimensions[column_letter].width = max(max_length + 2, 12)


def _looks_like_internal_key_row(row: tuple[Any, ...], dataset_name: str) -> bool:
    allowed = set(_headers_for_dataset(dataset_name, DATASET_SAMPLE_ROWS.get(dataset_name, []))) | set(FIELD_LABELS)
    values = [_cell_to_text(cell) for cell in row if _cell_to_text(cell)]
    return bool(values) and sum(1 for value in values if value in allowed) >= max(1, len(values) // 2)


def _resolve_field_key(dataset_name: str, header: str) -> str:
    cleaned = header.removesuffix("*").strip()
    if cleaned in FIELD_LABELS:
        return cleaned
    for key in _headers_for_dataset(dataset_name, DATASET_SAMPLE_ROWS.get(dataset_name, [])):
        if cleaned == FIELD_LABELS.get(key) or cleaned == key:
            return key
    for key, label in FIELD_LABELS.items():
        if cleaned == label:
            return key
    return cleaned


def _internal_value(key: str, value: str) -> str:
    choices = FIELD_CHOICES.get(key)
    if not choices:
        return value
    return choices.get(value, value)


def _display_value(key: str, value: Any) -> str:
    text = _cell_to_text(value)
    choices = FIELD_CHOICES.get(key)
    if not choices:
        return text
    reverse_choices = {internal: label for label, internal in choices.items()}
    return reverse_choices.get(text, text)


def _display_row(dataset_name: str, row: dict[str, str]) -> dict[str, str]:
    fields = _field_specs_for_dataset(dataset_name, [row])
    return {field.label: _display_value(field.key, row.get(field.key, "")) for field in fields}


def _missing_required_fields(dataset_name: str, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    required = DATASET_REQUIRED_FIELDS.get(dataset_name, set())
    missing: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=DATA_START_ROW):
        missing.extend(
            {"row": row_index, "field": key, "label": FIELD_LABELS.get(key, key)}
            for key in required
            if not _cell_to_text(row.get(key))
        )
    return missing


def _fallback_label(_key: str) -> str:
    return "扩展字段"


def _read_package_dataset(package_root: Path, dataset_name: str) -> list[dict[str, str]]:
    relative_path = DATASET_PATHS[dataset_name]
    path = package_root / relative_path
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return [dict(row) for row in csv.DictReader(file)]


def _sanitize_row(dataset_name: str, row: Mapping[str, Any]) -> dict[str, str]:
    sanitized = {key: _sanitize_value(key, value) for key, value in row.items()}
    if "enterprise_code" in sanitized:
        sanitized["enterprise_code"] = SANITIZED_ENTERPRISE_CODE
    if "tenant_slug" in sanitized:
        sanitized["tenant_slug"] = SANITIZED_TENANT_SLUG
    if "enterprise_name" in sanitized:
        sanitized["enterprise_name"] = SANITIZED_ENTERPRISE_NAME
    if dataset_name == "reporting_organizations" and "name" in sanitized:
        sanitized["name"] = SANITIZED_ENTERPRISE_NAME
    if sanitized.get("source_name"):
        sanitized["source_name"] = SANITIZED_SOURCE_NAME
    for path_key in ("attachment", "evidence_path", "package_root", "source_ref", "relative_path"):
        if path_key in sanitized:
            sanitized[path_key] = ""
    return sanitized


def _sanitize_value(key: str, value: Any) -> str:
    text = _cell_to_text(value)
    if key in {"remark", "note", "description"} and not text:
        return text
    replacements = {
        "本田金属技术(佛山)有限公司": SANITIZED_ENTERPRISE_NAME,
        f"本田金属技术{chr(0xFF08)}佛山{chr(0xFF09)}有限公司": SANITIZED_ENTERPRISE_NAME,
        "本田金属": "样例金属",
        "fs-honda-metal-2025": SANITIZED_TENANT_SLUG,
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()
