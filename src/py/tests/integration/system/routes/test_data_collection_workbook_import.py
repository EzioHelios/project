from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.domain.carbon.data_collection_workbook import (
    SANITIZED_ENTERPRISE_CODE,
    SANITIZED_ENTERPRISE_NAME,
    SANITIZED_TENANT_SLUG,
    build_collection_workbook_from_package,
    parse_data_collection_workbook,
)

if TYPE_CHECKING:
    from httpx import AsyncClient
    from sqlalchemy.ext.asyncio import AsyncSession


pytestmark = pytest.mark.anyio

HONDA_PACKAGE_DIR = (
    Path(__file__).resolve().parents[6] / "docs/评价方案/real_company_standardized_data/enterprises/honda"
)


async def test_sanitized_data_collection_workbook_imports_honda_energy_carbon_data(  # noqa: PLR0915
    client: AsyncClient,
    session: AsyncSession,
    superuser_token_headers: dict[str, str],
    tmp_path: Path,
) -> None:
    workbook_path = build_collection_workbook_from_package(
        HONDA_PACKAGE_DIR,
        tmp_path / "honda-metal-sanitized-data-collection.xlsx",
    )
    workbook_bytes = workbook_path.read_bytes()
    template_response = await client.get(
        "/api/carbon/data-collection-workbooks/template",
        headers=superuser_token_headers,
    )
    assert template_response.status_code == 200
    template_workbook = load_workbook(BytesIO(template_response.content), read_only=True, data_only=True)
    template_sheet = template_workbook["08_Scope1移动源燃烧"]
    template_visible_headers = [cell.value for cell in template_sheet[1]]
    template_internal_headers = [cell.value for cell in template_sheet[2]]
    assert "记录日期*" in template_visible_headers
    assert "燃料类型*" in template_visible_headers
    assert "record_date" in template_internal_headers
    assert "fuel_type" in template_internal_headers

    source_workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    source_sheet = source_workbook["08_Scope1移动源燃烧"]
    visible_headers = [cell.value for cell in source_sheet[1]]
    internal_headers = [cell.value for cell in source_sheet[2]]
    assert "记录日期*" in visible_headers
    assert "燃料类型*" in visible_headers
    assert "record_date" not in visible_headers
    assert "fuel_type" not in visible_headers
    assert "record_date" in internal_headers
    assert "fuel_type" in internal_headers

    datasets = parse_data_collection_workbook(workbook_bytes)
    assert datasets["enterprise_profiles"][0]["enterprise_code"] == SANITIZED_ENTERPRISE_CODE
    assert datasets["enterprise_profiles"][0]["tenant_slug"] == SANITIZED_TENANT_SLUG
    assert datasets["enterprise_profiles"][0]["enterprise_name"] == SANITIZED_ENTERPRISE_NAME
    assert len(datasets["energy_interval_facts"]) == 276
    assert len(datasets["scope1_mobile_combustion"]) == 1
    assert len(datasets["scope1_stationary_combustion"]) == 60
    assert len(datasets["scope1_refrigerant_leak"]) == 1
    assert len(datasets["scope2_electricity_bills"]) == 13
    assert len(datasets["scope2_purchased_energy"]) == 0
    assert len(datasets["scope3_third_party_transport"]) == 1
    assert len(datasets["product_footprint_inventory_items"]) == 10
    _append_scope2_purchased_energy_sample(workbook_path)
    workbook_bytes = workbook_path.read_bytes()
    datasets = parse_data_collection_workbook(workbook_bytes)
    assert len(datasets["scope2_purchased_energy"]) == 1
    assert datasets["scope2_purchased_energy"][0]["energy_type_code"] == "purchased_steam"

    preview_response = await client.post(
        "/api/carbon/data-collection-workbooks/preview",
        headers=superuser_token_headers,
        json={"workbook_base64": base64.b64encode(workbook_bytes).decode("ascii")},
    )

    assert preview_response.status_code == 200
    preview_payload = preview_response.json()["data"]
    assert preview_payload["total_rows"] >= 1
    purchased_energy_preview = next(
        item for item in preview_payload["sheets"] if item["dataset_name"] == "scope2_purchased_energy"
    )
    assert purchased_energy_preview["row_count"] == 1
    assert purchased_energy_preview["missing_required"] == []
    assert purchased_energy_preview["preview_rows"][0]["能源品种"] == "外购蒸汽"

    response = await client.post(
        "/api/carbon/data-collection-workbooks/import",
        headers=superuser_token_headers,
        json={"workbook_base64": base64.b64encode(workbook_bytes).decode("ascii")},
    )

    assert response.status_code == 201
    payload = response.json()["data"]
    assert payload["enterprise_code"] == SANITIZED_ENTERPRISE_CODE
    assert payload["tenant_slug"] == SANITIZED_TENANT_SLUG
    assert payload["imported"]["energy_interval_facts"] == 276
    assert payload["imported"]["scope1_mobile_combustion"] == 1
    assert payload["imported"]["scope1_stationary_combustion"] == 60
    assert payload["imported"]["scope1_refrigerant_leak"] == 1
    assert payload["imported"]["scope2_electricity_bills"] == 13
    assert payload["imported"]["scope2_purchased_energy"] == 1
    assert payload["imported"]["scope3_waste_disposal"] == 3
    assert payload["imported"]["scope3_third_party_transport"] == 1
    assert payload["imported"]["production_output_facts"] == 24
    assert payload["imported"]["product_footprint_inventory_items"] == 10
    assert payload["imported"]["product_footprint_results"] == 1

    team_id = payload["team_id"]
    assert await _count(session, "team", team_id, id_column="id") == 1
    assert await _count(session, "dataapp_energyintervalfact", team_id) == 276
    assert await _count(session, "dataapp_productionoutputfact", team_id) == 24
    assert await _count(session, "dataapp_scope1mobilecombustion", team_id) == 1
    assert await _count(session, "dataapp_scope1stationarycombustion", team_id) == 60
    assert await _count(session, "dataapp_scope1refrigerantleak", team_id) == 1
    assert await _count(session, "dataapp_scope2electricitybill", team_id) == 13
    assert await _count(session, "dataapp_scope2purchasedenergy", team_id) == 1
    assert await _count(session, "dataapp_scope3wastedisposal", team_id) == 3
    assert await _count(session, "dataapp_scope3thirdpartytransport", team_id) == 1
    assert await _count(session, "carbon_product_footprint_inventory_item", team_id) == 10
    assert await _count(session, "carbon_product_footprint_result", team_id) == 1

    team_name = await session.scalar(text("select name from team where id = :team_id"), {"team_id": team_id})
    assert team_name == SANITIZED_ENTERPRISE_NAME
    local_path_count = await session.scalar(
        text(
            """
            select count(*)
            from dataapp_scope1stationarycombustion
            where team_id = :team_id and coalesce(attachment, '') like '%/Users/%'
            """
        ),
        {"team_id": team_id},
    )
    assert local_path_count == 0


def _append_scope2_purchased_energy_sample(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook["12_Scope2外购热汽冷"]
    internal_headers = [cell.value for cell in sheet[2]]
    row = {
        "energy_type_code": "外购蒸汽",
        "energy_type_name": "外购蒸汽",
        "supplier_name": "脱敏园区供热站",
        "period_start": "2025-01-01",
        "period_end": "2025-01-31",
        "amount": "72.0000",
        "unit": "t",
        "standard_amount": "9259.2000",
        "standard_unit": "kgce",
        "tce_amount": "9.2592",
        "amount_cny": "17640.00",
        "area_id": "AREA-SAMPLE",
        "area_name": "脱敏样例区域",
        "facility_name": "脱敏换热站",
        "remark": "脱敏补充样例, 用于验证外购蒸汽导入链路.",
        "calc_status": "ready",
        "factor_source": "脱敏蒸汽排放因子",
        "factor_year": "2025",
        "factor_unit": "kgCO2e/t",
        "activity_amount": "72.0000",
        "activity_unit": "t",
        "computed_co2e": "7200.000000",
        "computation_note": "脱敏补充样例, 不代表本田源数据.",
        "standard_code": "GB/T 2589-2020",
        "conversion_factor_source": "脱敏样例折标口径.",
    }
    sheet.append([row.get(str(header), "") for header in internal_headers])
    workbook.save(workbook_path)


async def _count(session: AsyncSession, table_name: str, team_id: str, *, id_column: str = "team_id") -> int:
    value = await session.scalar(
        text(f"select count(*) from {table_name} where {id_column} = :team_id"),
        {"team_id": team_id},
    )
    return int(value or 0)
