# mypy: disable-error-code=import-untyped

import base64
from io import BytesIO

from openpyxl import load_workbook

from app.domain.ghg_inventory.infrastructure.workbook import build_template_bytes, decode_and_preview


def _context(method_code: str = "F08_PURCHASED_ENERGY", category_code: str = "2.1") -> dict[str, str]:
    return {
        "inventory_code": "INV-2026",
        "organization_code": "ORG-01",
        "facility_code": "FAC-01",
        "iso_category_code": category_code,
        "reporting_profile_code": "CORPORATE_ISO14064_GHGP",
        "method_code": method_code,
    }


def _append_by_code(
    workbook: object,
    sheet_name: str,
    values: dict[str, object],
    *,
    row_number: int | None = None,
) -> None:
    sheet = workbook[sheet_name]  # type: ignore[index]
    codes = [str(cell.value or "") for cell in sheet[2]]
    if row_number is None:
        sheet.append([values.get(code) for code in codes])
        return
    for column_number, code in enumerate(codes, 1):
        if code in values:
            sheet.cell(row=row_number, column=column_number, value=values[code])


def _valid_purchased_electricity_workbook() -> tuple[str, dict[str, str]]:
    context = _context()
    workbook = load_workbook(BytesIO(build_template_bytes(context)))
    _append_by_code(
        workbook,
        "01_组织与边界",
        {
            "inventory_code": context["inventory_code"],
            "organization_code": context["organization_code"],
            "facility_code": context["facility_code"],
            "reporting_profile_code": context["reporting_profile_code"],
            "iso_category_code": context["iso_category_code"],
            "boundary_method": "operational_control",
        },
        row_number=3,
    )
    _append_by_code(
        workbook,
        "02_排放源",
        {
            "source_code": "PCC-01",
            "source_name": "购入电力关口",
            "activity_type": "purchased_electricity",
            "facility_code": context["facility_code"],
            "iso_category_code": context["iso_category_code"],
        },
    )
    _append_by_code(
        workbook,
        "03_活动数据",
        {
            "record_code": "ELEC-001",
            "source_code": "PCC-01",
            "activity_code": "PURCHASED-ELECTRICITY",
            "activity_name": "购入电力",
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "activity_amount": 1000,
            "activity_unit": "kWh",
            "data_source_type": "measured",
            "measurement_method": "关口电表",
            "parameter_set_code": "EF-JS-2026",
            "boundary_status": "included",
            "energy_type": "electricity",
            "supplier_name": "电网企业",
            "meter_point": "PCC-01",
            "scope2_method": "location",
        },
    )
    _append_by_code(
        workbook,
        "03A_区域设备归集",
        {
            "allocation_code": "ELEC-001-A01",
            "record_code": "ELEC-001",
            "source_code": "PCC-01",
            "area_code": "ROOT",
            "measurement_boundary_code": "PCC-01",
            "allocation_role": "energy_consumption",
            "allocation_ratio": 1,
            "source_kind": "manual",
        },
    )
    _append_by_code(
        workbook,
        "04_参数与因子",
        {
            "parameter_set_code": "EF-JS-2026",
            "parameter_code": "ENERGY_FACTOR",
            "parameter_name": "电力排放因子",
            "parameter_role": "emission_factor",
            "value": 0.55,
            "unit": "kgCO2e/kWh",
            "source_mode": "authority_published",
            "source_name": "主管部门公告",
            "source_version": "2026",
            "standard_code": "GB/T 32150-2025",
            "evidence_code": "EVID-EF-001",
        },
        row_number=3,
    )
    for evidence_code, file_name, document_type in (
        ("EVID-EF-001", "排放因子公告.pdf", "standard"),
        ("EVID-METER-001", "电表台账.xlsx", "meter_ledger"),
    ):
        _append_by_code(
            workbook,
            "05_证据索引",
            {
                "evidence_code": evidence_code,
                "file_name": file_name,
                "document_type": document_type,
                "storage_key": f"evidence/{file_name}",
                "sha256": "a" * 64 if evidence_code.endswith("001") else "b" * 64,
                "confidentiality": "internal",
                "verification_status": "verified",
            },
        )
    _append_by_code(
        workbook,
        "06_数据证据关联",
        {
            "link_code": "LINK-ACT-001",
            "record_type": "activity",
            "record_code": "ELEC-001",
            "evidence_code": "EVID-METER-001",
            "evidence_role": "activity",
            "coverage_ratio": 1,
        },
    )
    stream = BytesIO()
    workbook.save(stream)
    return base64.b64encode(stream.getvalue()).decode(), context


def test_template_has_stable_codes_and_bounded_visible_sheets() -> None:
    workbook = load_workbook(BytesIO(build_template_bytes(_context())))
    assert workbook["_metadata"].sheet_state == "veryHidden"
    assert workbook["03_活动数据"].row_dimensions[2].hidden is True
    assert workbook["03_活动数据"].freeze_panes == "A3"
    assert sum(sheet.sheet_state == "visible" for sheet in workbook.worksheets) == 8


def test_valid_workbook_passes_precheck() -> None:
    encoded, context = _valid_purchased_electricity_workbook()
    preview, _file_hash = decode_and_preview(encoded, context)
    assert preview.payload()["metrics"] == {
        "rowCount": 1,
        "validCount": 1,
        "warningCount": 0,
        "errorCount": 0,
        "missingEvidenceCount": 0,
    }


def test_parameter_unit_and_duplicate_are_blocked() -> None:
    encoded, context = _valid_purchased_electricity_workbook()
    workbook = load_workbook(BytesIO(base64.b64decode(encoded)))
    parameter_sheet = workbook["04_参数与因子"]
    codes = [str(cell.value or "") for cell in parameter_sheet[2]]
    unit_column = codes.index("unit") + 1
    parameter_sheet.cell(row=3, column=unit_column, value="kgCO2e/km")
    parameter_sheet.append([cell.value for cell in parameter_sheet[3]])
    stream = BytesIO()
    workbook.save(stream)

    preview, _file_hash = decode_and_preview(base64.b64encode(stream.getvalue()).decode(), context)
    issue_codes = {issue.issue_code for issue in preview.issues}
    assert "PARAMETER_UNIT_NOT_ALLOWED" in issue_codes
    assert "DUPLICATE_PARAMETER" in issue_codes


def test_non_executable_method_is_blocked_at_precheck() -> None:
    context = _context(method_code="F10_VALUE_CHAIN", category_code="4.1")
    workbook = load_workbook(BytesIO(build_template_bytes(context)))
    _append_by_code(
        workbook,
        "01_组织与边界",
        {
            "inventory_code": context["inventory_code"],
            "organization_code": context["organization_code"],
            "facility_code": context["facility_code"],
            "reporting_profile_code": context["reporting_profile_code"],
            "iso_category_code": context["iso_category_code"],
            "boundary_method": "operational_control",
        },
        row_number=3,
    )
    _append_by_code(
        workbook,
        "02_排放源",
        {
            "source_code": "PROCUREMENT",
            "source_name": "采购活动",
            "activity_type": "purchased_goods",
            "facility_code": context["facility_code"],
            "iso_category_code": context["iso_category_code"],
        },
    )
    _append_by_code(
        workbook,
        "03_活动数据",
        {
            "record_code": "PURCHASE-001",
            "source_code": "PROCUREMENT",
            "activity_code": "MATERIAL",
            "activity_name": "采购材料",
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "activity_amount": 100,
            "activity_unit": "kg",
            "data_source_type": "ledger",
            "measurement_method": "采购台账",
            "parameter_set_code": "PCF-001",
            "boundary_status": "included",
            "supplier_name": "供应商",
            "product_or_service": "材料",
            "lifecycle_boundary": "cradle-to-gate",
        },
    )
    _append_by_code(
        workbook,
        "03A_区域设备归集",
        {
            "allocation_code": "PURCHASE-001-A01",
            "record_code": "PURCHASE-001",
            "source_code": "PROCUREMENT",
            "area_code": "ROOT",
            "allocation_role": "beneficiary",
            "allocation_ratio": 1,
            "external_facility_code": "SUPPLIER-01",
            "source_kind": "manual",
        },
    )
    _append_by_code(
        workbook,
        "04_参数与因子",
        {
            "parameter_set_code": "PCF-001",
            "parameter_code": "VALUE_CHAIN_FACTOR",
            "parameter_name": "供应商产品碳足迹",
            "parameter_role": "emission_factor",
            "value": 2.1,
            "unit": "kgCO2e/unit",
            "source_mode": "supplier",
            "source_name": "供应商声明",
        },
        row_number=3,
    )
    stream = BytesIO()
    workbook.save(stream)
    preview, _file_hash = decode_and_preview(base64.b64encode(stream.getvalue()).decode(), context)
    assert "METHOD_NOT_EXECUTABLE" in {issue.issue_code for issue in preview.issues}


def test_evidence_links_validate_source_and_boundary_references() -> None:
    encoded, context = _valid_purchased_electricity_workbook()
    workbook = load_workbook(BytesIO(base64.b64decode(encoded)))
    _append_by_code(
        workbook,
        "06_数据证据关联",
        {
            "link_code": "LINK-SOURCE-INVALID",
            "record_type": "source",
            "record_code": "UNKNOWN-SOURCE",
            "evidence_code": "EVID-METER-001",
            "evidence_role": "activity",
        },
    )
    _append_by_code(
        workbook,
        "06_数据证据关联",
        {
            "link_code": "LINK-BOUNDARY-INVALID",
            "record_type": "boundary",
            "record_code": "UNKNOWN-INVENTORY",
            "evidence_code": "EVID-METER-001",
            "evidence_role": "boundary",
        },
    )
    stream = BytesIO()
    workbook.save(stream)

    preview, _file_hash = decode_and_preview(base64.b64encode(stream.getvalue()).decode(), context)
    issue_codes = {issue.issue_code for issue in preview.issues}
    assert "SOURCE_NOT_FOUND" in issue_codes
    assert "BOUNDARY_NOT_FOUND" in issue_codes
