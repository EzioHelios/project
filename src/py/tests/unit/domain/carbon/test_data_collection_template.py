from io import BytesIO

from openpyxl import load_workbook

from app.domain.carbon.data_collection_workbook import build_data_collection_template_bytes


def test_data_collection_template_has_no_generated_business_rows_or_development_copy() -> None:
    workbook = load_workbook(BytesIO(build_data_collection_template_bytes()), read_only=True, data_only=True)
    metadata_text = " ".join(
        str(value) for row in workbook["00_填报说明"].iter_rows(values_only=True) for value in row if value is not None
    )

    assert "样例" not in metadata_text
    assert "脱敏" not in metadata_text
    assert "联调" not in metadata_text
    for sheet in workbook.worksheets:
        if sheet.title != "00_填报说明":
            assert not any(
                any(value not in (None, "") for value in row) for row in sheet.iter_rows(min_row=3, values_only=True)
            )
