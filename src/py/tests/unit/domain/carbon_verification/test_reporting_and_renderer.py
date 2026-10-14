from __future__ import annotations

import zipfile
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from typing import TYPE_CHECKING
from xml.etree import ElementTree as ET

import pytest

from app.domain.carbon_verification.domain.regulatory_templates import select_descriptor
from app.domain.carbon_verification.domain.reporting import formal_report_readiness
from app.domain.carbon_verification.infrastructure.report_renderer import (
    ReportDocumentRenderer,
    _build_document,
    _device_month_groups,
)

if TYPE_CHECKING:
    from pathlib import Path


def _report() -> dict[str, object]:
    activity = {
        "record_code": "ACT-001",
        "scope": "scope2",
        "iso_category_code": "2",
        "activity_name": "外购电力",
        "activity_amount": 1000,
        "activity_unit": "kWh",
        "computed_co2e": 570,
        "evidence_status": "ready",
    }
    return {
        "project": {
            "report_code": "RPT-2026-001",
            "report_version": 1,
            "run_code": "RUN-2026-001",
            "verification_basis": "ISO 14064-3:2019",
        },
        "template": {
            "template_code": "organization_ghg_generic",
            "standard_name": "工业企业温室气体排放核算和报告通则",
            "standard_code": "GB/T 32150—2025",
            "version": "v2026.07.01",
            "schema_version": "2026.1",
            "implementation_status": "supported",
            "renderer_kind": "generic_a4",
            "report_purpose": "organization_inventory",
        },
        "standard": {"name": "ISO 14064-1", "code": "ISO 14064-1:2018"},
        "enterprise": {"name": "示例制造企业", "code": "ORG-001", "unify_code": "91320000EXAMPLE"},
        "reporting_period": {
            "year": 2026,
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
            "base_year": 2025,
        },
        "inventory_source": {
            "status": "published",
            "boundary_no": "BOUNDARY-2026-01",
            "boundary_status": "published",
        },
        "inventory": {
            "totals": {"total_kg_co2e": 570, "total_t_co2e": 0.57},
            "scope_summary": [{"scope": "scope2", "kg_co2e": 570, "t_co2e": 0.57}],
            "source_summary": [{"scope": "scope2", "record_label": "外购电力", "record_count": 1, "t_co2e": 0.57}],
            "gas_summary": [{"gas_code": "co2", "kg_co2e": 570, "t_co2e": 0.57}],
            "iso_category_summary": [{"label": "输入能源产生的间接温室气体排放", "kg_co2e": 570, "t_co2e": 0.57}],
            "scope2_method_summary": [{"method": "location_based", "kg_co2e": 570, "t_co2e": 0.57}],
            "scope3_category_summary": [],
            "gwp_versions": ["IPCC AR6"],
            "default_scope2_method": "location_based",
        },
        "activity_data": [activity],
        "factor_parameters": [
            {
                "record_code": "ACT-001",
                "scope": "scope2",
                "factor_source": "全国电网排放因子",
                "factor_version": "2026",
                "factor_unit": "kgCO2e/kWh",
            }
        ],
        "verification_attachments": [
            {
                "evidence_code": "EV-001",
                "record_code": "ACT-001",
                "file_name": "电费账单.pdf",
                "attachment_type": "invoice",
                "material_status": "ready",
                "sha256": "a" * 64,
            }
        ],
        "quality_issues": [],
        "verification_findings": [],
        "exclusions": [],
        "formal_readiness": {"status": "ready", "blocking_count": 0, "warning_count": 0, "checks": []},
    }


def test_factor_reference_does_not_replace_activity_evidence():
    report = _report()
    report["verification_attachments"][0]["evidence_role"] = "factor"
    checks = {item.code: item.status for item in formal_report_readiness(report).checks}
    assert checks["evidence"] == "blocked"


def test_quality_summary_and_parameter_precision():
    report = _report()
    report["factor_parameters"][0].update(
        parameter_name="购入能源排放因子", parameter_code="ENERGY_FACTOR", factor_value="0.582700000000"
    )
    report["quality_issues"] = [
        {"record_code": "ACT-001", "title": "核算记录待复核", "action": "复核资料"} for _ in range(3)
    ]
    doc = _build_document(report)
    texts = [paragraph.text for paragraph in doc.paragraphs]
    texts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    text = "\n".join(texts)
    assert text.count("核算记录待复核") == 1
    assert "0.5827" in text and "0.582700000000" not in text
    assert "ENERGY_FACTOR" not in text


@pytest.mark.unit
def test_formal_report_readiness_requires_a_complete_frozen_chain() -> None:
    report = _report()
    readiness = formal_report_readiness(report)
    assert readiness.status == "ready"
    report["verification_attachments"] = []
    blocked = formal_report_readiness(report)
    assert blocked.status == "blocked"
    assert any(check.code == "evidence" and check.status == "blocked" for check in blocked.checks)


@pytest.mark.unit
def test_report_renderer_creates_a4_docx_with_dynamic_rows(tmp_path: Path) -> None:
    report = _report()
    report["formal_readiness"] = {
        "status": "ready",
        "blocking_count": 0,
        "warning_count": 0,
        "checks": [],
    }
    files = ReportDocumentRenderer(template_path=None, output_dir=tmp_path).render(report)
    docx = next(item for item in files if item.path.endswith(".docx"))
    with zipfile.ZipFile(BytesIO(docx.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        section_xml = document_xml[document_xml.rfind("<w:sectPr") :]
    assert "示例制造企业" in document_xml
    assert "ACT-001" not in document_xml
    assert "参数及活动记录对应关系" not in document_xml
    assert "RPT-2026-001" not in document_xml
    assert "RUN-2026-001" not in document_xml
    assert "EV-001" not in document_xml
    assert "电费账单.pdf" in document_xml
    assert 'w:w="11906"' in section_xml
    assert 'w:h="16838"' in section_xml


@pytest.mark.unit
def test_report_renderer_rejects_unimplemented_regulatory_adapter(tmp_path: Path) -> None:
    report = _report()
    template = report["template"]
    assert isinstance(template, dict)
    template["implementation_status"] = "registered"
    template["renderer_kind"] = "ets_power_2026"
    with pytest.raises(ValueError, match="尚未实现"):
        ReportDocumentRenderer(template_path=None, output_dir=tmp_path).render(report)


@pytest.mark.unit
def test_national_ets_routing_never_silently_falls_back_to_generic() -> None:
    descriptor = select_descriptor(purpose="national_ets", industry="钢铁制造", year=2026)
    assert descriptor.standard_code == "CETS-AG-03.01-V01-2024"
    with pytest.raises(ValueError, match="未找到"):
        select_descriptor(purpose="national_ets", industry="未知行业", year=2026)


@pytest.mark.unit
@pytest.mark.parametrize(
    ("status", "label"),
    [
        ("draft", "草稿"),
        ("submitted", "待审核"),
        ("approved", "已审核"),
        ("rejected", "已退回"),
        ("published", "已发布"),
        ("archived", "已归档"),
        ("", "【待确认】"),
    ],
)
def test_report_identity_period_and_lifecycle_are_not_fabricated(status: str, label: str) -> None:
    report = _report()
    report["enterprise"] = {
        "name": "祥泰仓库",
        "code": "INTERNAL-NOT-USCC",
        "unify_code": None,
        "organization_type": "facility",
    }
    report["inventory_source"] = {"status": status, "boundary_name": "已接入设备计量范围"}
    report["reporting_period"] = {
        "year": 2026,
        "period_start": "2026-01-01",
        "period_end": "2026-09-30",
        "data_period_start": "2026-02-01",
        "data_period_end": "2026-08-31",
    }
    doc = _build_document(report)
    text = " ".join(ET.fromstring(doc.element.xml).itertext())
    assert "INTERNAL-NOT-USCC" not in text
    assert "【待补充" in text
    assert "祥泰仓库" in text
    assert "设施" in text
    assert "2026-01-01 至 2026-09-30" in text
    assert "2026-02-01 至 2026-08-31" in text
    assert "2026-12-31" not in text
    assert doc.tables[0].cell(2, 0).text == "报告状态"
    assert doc.tables[0].cell(2, 1).text == label
    if status != "published":
        assert "已发布" not in text
    assert "发布时冻结" not in text


@pytest.mark.unit
def test_summary_page_break_is_attached_to_heading_not_trailing_empty_paragraph() -> None:
    doc = _build_document(_report())
    heading = next(p for p in doc.paragraphs if p.text == "四、温室气体排放结果" and p.style.name == "Heading 1")
    assert heading.paragraph_format.page_break_before is True
    previous = heading._p.getprevious()
    assert not previous.xpath('.//w:br[@w:type="page"]')
    assert not previous.xpath(".//w:pageBreakBefore")


def test_device_month_aggregation_preserves_identity_units_zero_and_calculation_shares():
    activity = _report()["activity_data"][0]
    activity.update(
        period_start="2026-01-01",
        period_end="2026-01-31",
        device_allocations=[
            {"device_id": 1, "device_name": "同名设备", "ratio": "0.25"},
            {"device_id": 2, "device_name": "同名设备", "ratio": "0.75"},
        ],
    )
    february = deepcopy(activity)
    february.update(
        record_code="ACT-002", period_start="2026-02-01", period_end="2026-02-28", activity_amount=0, computed_co2e=0
    )
    fuel = deepcopy(activity)
    fuel.update(record_code="ACT-003", activity_unit="L", iso_category_code="1.1")
    groups = _device_month_groups([activity, february, fuel])
    electric = groups[("2", "kWh")]
    assert len(electric) == 2
    assert electric[0]["months"] == {"2026-01": Decimal(250), "2026-02": Decimal(0)}
    assert electric[1]["months"]["2026-01"] == Decimal(750)
    assert sum(row["co2e_kg"] for row in electric) == Decimal(570)
    assert ("1.1", "L") in groups


def test_device_month_aggregation_does_not_spread_multi_month_records():
    activity = _report()["activity_data"][0]
    activity.update(period_start="2026-01-01", period_end="2026-03-31")
    grouped = _device_month_groups([activity])[("2", "kWh")][0]
    assert grouped["months"] == {"2026-01-01 至 2026-03-31": Decimal(1000)}
    assert grouped["name"] == "【待补充：设备归属】"  # noqa: RUF001


@pytest.mark.parametrize("ratios", [("0.6", "0.6"), ("0.1", "0.8"), ("-0.2", "1.2")])
def test_device_month_aggregation_rejects_invalid_allocation(ratios):
    activity = _report()["activity_data"][0]
    activity["device_allocations"] = [{"device_id": i, "ratio": ratio} for i, ratio in enumerate(ratios, 1)]
    with pytest.raises(ValueError, match="归集比例"):
        _device_month_groups([activity])


def test_report_has_explicit_reference_fonts_and_navigation():
    doc = _build_document(_report())
    assert doc.styles["Normal"].font.size.pt == 14
    assert doc.styles["Heading 1"].font.size.pt == 15
    from docx.oxml.ns import qn

    fonts = doc.styles["Heading 1"]._element.rPr.rFonts
    assert fonts.get(qn("w:eastAsia")) == "宋体"
    assert fonts.get(qn("w:ascii")) == "Times New Roman"
    assert fonts.get(qn("w:asciiTheme")) is None
    assert len(doc.element.xpath(".//w:bookmarkStart")) == 8
    assert len(doc.element.xpath(".//w:fldSimple")) == 8
