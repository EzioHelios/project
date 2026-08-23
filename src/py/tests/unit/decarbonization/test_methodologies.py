"""Deterministic methodology and reduction-domain policy tests."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.domain.decarbonization.domain import (
    CcerIssuancePolicy,
    MethodologyInputError,
    ProjectLifecyclePolicy,
    ReductionDomainError,
    calculate_reduction,
    screen_applicability,
)
from app.domain.decarbonization.infrastructure.workbook import (
    build_monitoring_template,
    preview_monitoring_workbook,
)
from app.domain.decarbonization.usecases.projects import _level_irr

pytestmark = pytest.mark.unit


@pytest.mark.parametrize(
    ("code", "inputs", "expected"),
    [
        (
            "CCER-01-001-V01",
            {
                "exported_electricity_mwh": 100,
                "imported_electricity_mwh": 10,
                "grid_om_factor": 0.8,
                "grid_bm_factor": 0.6,
                "om_weight": 0.6,
                "fuel_emissions_tco2": 4,
            },
            Decimal("60.800"),
        ),
        (
            "CCER-01-004-V01",
            {
                "sold_hydrogen_t": 10,
                "renewable_power_mwh": 80,
                "grid_import_mwh": 20,
                "baseline_hydrogen_factor": 10,
            },
            Decimal("80.000"),
        ),
        (
            "CCER-06-001-V01",
            {
                "baseline_emissions_tco2e": 100,
                "project_emissions_tco2e": 70,
                "refrigerant_emissions_tco2e": 5,
                "minimum_monthly_use_hours": 160,
                "baseline_hdd": 100,
                "project_hdd": 110,
                "baseline_cdd": 200,
                "project_cdd": 190,
            },
            Decimal("25.000"),
        ),
        (
            "CCER-14-001-V01",
            {
                "biomass_change_tco2e": 100,
                "dead_organic_matter_change_tco2e": 10,
                "soil_carbon_change_tco2e": 20,
                "fire_emissions_tco2e": 5,
                "pre_existing_biomass_change_tco2e": 2,
                "risk_deduction_tco2e": 3,
            },
            Decimal("120.000"),
        ),
        (
            "CCER-14-002-V01",
            {
                "biomass_carbon_change_tc": 12,
                "soil_carbon_change_tc": 6,
                "soil_ch4_emissions_tco2e": 2,
                "soil_n2o_emissions_tco2e": 1,
                "risk_deduction_tco2e": 3,
            },
            Decimal("60.000"),
        ),
    ],
)
def test_executable_methodology_results_are_deterministic(
    code: str, inputs: dict[str, float], expected: Decimal
) -> None:
    result = calculate_reduction(code, inputs)

    assert result.reduction_tco2e == expected


def test_building_gate_sets_reduction_to_zero_when_weather_deviation_is_excessive() -> None:
    result = calculate_reduction(
        "CCER-06-001-V01",
        {
            "baseline_emissions_tco2e": 100,
            "project_emissions_tco2e": 50,
            "minimum_monthly_use_hours": 160,
            "baseline_hdd": 100,
            "project_hdd": 130,
        },
    )

    assert result.reduction_tco2e == 0
    assert any(item["level"] == "blocking" for item in result.diagnostics)


def test_methodology_rejects_reference_only_and_invalid_physical_inputs() -> None:
    with pytest.raises(MethodologyInputError, match="reference-only"):
        calculate_reduction("CCER-01-002-V01", {})
    with pytest.raises(MethodologyInputError, match="om_weight"):
        calculate_reduction(
            "CCER-01-001-V01",
            {
                "exported_electricity_mwh": 10,
                "imported_electricity_mwh": 0,
                "grid_om_factor": 0.8,
                "grid_bm_factor": 0.6,
                "om_weight": 1.2,
            },
        )


def test_applicability_requires_every_formal_check() -> None:
    result = screen_applicability(
        "CCER-01-004-V01",
        {"new_build": True, "same_legal_entity": True, "self_consumption": True, "platform_connected": False},
    )

    assert result["status"] == "not_applicable"
    assert result["blockingReasons"]


def test_project_lifecycle_and_ccer_issuance_boundaries() -> None:
    ProjectLifecyclePolicy.ensure_track("internal_action")
    ProjectLifecyclePolicy.ensure_transition("draft", "screening")
    with pytest.raises(ReductionDomainError):
        ProjectLifecyclePolicy.ensure_transition("draft", "issued")
    with pytest.raises(ReductionDomainError):
        CcerIssuancePolicy.ensure_asset_eligible(
            project_track="internal_action", dossier_status="issued", verified_tco2e=Decimal(100)
        )


def test_monitoring_workbook_round_trip_and_unit_gate() -> None:
    methodology = {
        "code": "CCER-01-004",
        "name": "可再生能源电解水制氢",
        "version": "V01",
        "source_url": "https://ccer.cets.org.cn/",
        "parameter_schema": [
            {
                "code": "sold_hydrogen_t",
                "name": "外售氢气量",
                "unit": "t",
                "minimum": 0,
                "required": True,
            }
        ],
        "evidence_schema": ["氢气计量记录"],
    }
    content = build_monitoring_template(methodology)
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["01_监测数据"]
    sheet.cell(2, 3, "2026-01-01")
    sheet.cell(2, 4, "2026-01-31")
    sheet.cell(2, 5, 12.5)
    sheet.cell(2, 7, "贸易结算表")
    sheet.cell(2, 8, "月度")
    sheet.cell(2, 9, "EV-001")
    stream = BytesIO()
    workbook.save(stream)

    valid = preview_monitoring_workbook(stream.getvalue(), methodology).payload()
    assert valid["canCommit"] is True
    assert valid["metrics"]["validCount"] == 1

    sheet.cell(2, 6, "kg")
    invalid_stream = BytesIO()
    workbook.save(invalid_stream)
    invalid = preview_monitoring_workbook(invalid_stream.getvalue(), methodology).payload()
    assert invalid["canCommit"] is False
    assert invalid["issues"][0]["issueCode"] == "unit_mismatch"


def test_level_cash_flow_irr_is_calculated_as_a_ratio() -> None:
    result = _level_irr(Decimal(100), Decimal(30), 5)

    assert result is not None
    assert result.quantize(Decimal("0.0001")) == Decimal("0.1524")
